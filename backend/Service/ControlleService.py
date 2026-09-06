import os
import random
import time
import math
import uuid
from datetime import datetime, timezone, timedelta
from threading import Thread, RLock
from queue import Queue

from skyfield.api import load

from Service.SatelliteNetworkService import SatelliteNetwork
from Service.TaskService import Task
from Service.exportfuc import exportcfuc
from Service.split import split_area_target_into_point_targets, split_cycle_target_into_point_targets
from config import INTER_VAL_TIME, TLE, sat_parms, cluster_parms
from database import db
from model.TaskModel import NewTaskModel, OldTaskModel
from utils.time_util import get_now_time_from_start
from flask import current_app
from contextlib import nullcontext
from skyfield.api import wgs84


def _local_naive_to_utc(dt):
    """
    前端/表格传入的 naive 时间按系统本地时区解释，统一转为 UTC naive。
    （naive.astimezone 在 Python3.6+ 按系统本地时区解释；转完保持 naive 交给下游，下游自行 replace(tzinfo=utc)）
    """
    if dt is None:
        return None
    return dt.astimezone(timezone.utc).replace(tzinfo=None)


class OperationsControlCenter:
    """
    控制中心，负责实时地从用户处接收任务，实时地获取卫星网络状态，并进行任务的规划，部署到卫星网络中，并返回结果给用户
    """

    # 初始化
    def __init__(self):
        """
        初始化控制中心
        """
        self.app = None  # 将在run方法中设置
        # 共享任务列表/字典（new_tasks、pause_tasks、net_tasks_buffer、tasks_buffer、appointment_tasks 等）的并发锁，
        # Flask请求线程、OCC规划线程、卫星网络线程三方共用同一把锁；用RLock防止持锁回调同锁函数时死锁
        self.task_lock = RLock()
        self.task_queue = Queue()  # 任务队列
        self.result_queue = Queue()  # 从数据中心接收结果
        self.network_state = None  # 卫星网络状态
        self.satellite_network = None  # 卫星网络对象
        self.data_center_queue = Queue()  # 发送到数据中心的队列

        self.system_interval = 5  # 系统时间间隔，默认为5秒，单位：秒

        # self.start_time = datetime.strptime("2021-05-10 00:00:00", "%Y-%m-%d %H:%M:%S")
        # self.end_time = datetime.strptime("2021-05-16 00:00:00", "%Y-%m-%d %H:%M:%S")
        self.start_time = datetime.strptime("2025-06-06 00:00:00", "%Y-%m-%d %H:%M:%S")
        self.end_time = datetime.strptime("2025-06-13 00:00:00", "%Y-%m-%d %H:%M:%S")
        self.now_time = self.start_time  # 当前时间
        self.time_multiple = 1  # 时间倍数,默认为1倍，即正常速度
        self.tasks_buffer = {}  # 任务缓冲区
        self.appointment_tasks = []  # 定时任务列表

        # 算法参数
        self.model = 0  # 0代表最优方案，1代表最小化方案，2代表资源利用率最大方案，3代表成像质量最高方案
        self.completed_gravity = 1.0  # 完成度权重
        self.balance_gravity = 1.0  # 平衡权重
        self.priority_gravity = 1.0  # 优先级权重
        self.auto_run = True  # 运行模式：True=自主运行（系统自动规划），False=程序控制（手动触发规划）
        # 星簇级约束项自定义管理（时间/能源/固存三类），阈值为 0 表示不额外收紧，在 exportfuc 规划输入装配处统一生效
        self.constraint_config = {
            'time': {'name': '时间约束', 'enabled': True, 'threshold': 0.0, 'unit': 'min',
                     'desc': '任务时间窗两端各收缩的缓冲时长'},
            'energy': {'name': '能源约束', 'enabled': True, 'threshold': 0.0, 'unit': 'Wh',
                       'desc': '每颗卫星预留的最低剩余电量（调度时不可用）'},
            'storage': {'name': '固存约束', 'enabled': True, 'threshold': 0.0, 'unit': 'GB',
                        'desc': '每颗卫星预留的最低剩余存储（调度时不可用）'},
        }
        self.comprehensive_case_ids = None  # 综合验证案例已生成的任务ID列表（本会话内防重复生成）
        self.statistical_data = []  # 统计数据
        self.cluster_data = []  # 星簇数据
        self.planning_results = []  # 规划结果

        # 是否使用随机化方式初始化卫星
        self.is_random = False  # 是否使用随机化方式初始化卫星
        # 是否提交了卫星参数文件
        self.is_submit_sat = False
        # 是否使用生成程序来产生TLE
        self.is_TLE = False
        # 是否提交了TLE文件
        self.is_submit_tle = False
        # 是否提交了系统参数
        self.is_submit_sys = False
        # 卫星网络初始化是否失败（OCC 线程构造 SatelliteNetwork 异常时置 True）
        self.init_failed = False
        # 最新批次规划任务使用的算法名字
        self.algorithm_name = None  # 最新批次规划任务使用的算法名字
        # 默认倍速
        self.default_speed_doubling = 20  # 默认倍速

        # 三种案例最新的id
        self.point = None
        self.area = None
        self.ocean = None

        self.sats_config = {}

        # 卫星网络配置
        self.network_config = {
            'Sat_Walker': {
                'TLE_path': f'./library/{TLE}',  # TLE文件路径
                'Satellite_Params_Path': f'./library/{sat_parms}',  # 卫星的配置文件
                'Clusters_path': f'./library/{cluster_parms}',  # 集群的配置文件
                'Num_of_Orbits': 10,  # 轨道数
                'Sat_Num_per_Orbit': 20,  # 每轨道卫星数
                'Start_Time': self.start_time,  # 仿真开始时间
                'End_Time': self.end_time,  # 仿真结束时间
                'interval_time': 5 * 60  # 秒，仿真间隔时间
            },
            'SAR': {
                'max_storage': 500,
                'max_battery': 5000,
                'resolution': 0.5,
                'pitch_angle': 45,
                'side_swing_angle': 45,
                'stable_time': 10,
                'angle_velocity': 5,
                'width_of_cloth': 50,
                'thickness_threshold': 800,
                'imaging_powers': 50,
                'sunlight_powers': 30,
                'eclipse_powers': 8,
                'maneuver_powers': 500,
                'downlink_rate': 4
            },
            'optical': {
                'max_storage': 500,
                'max_battery': 5000,
                'resolution': 0.5,
                'pitch_angle': 45,
                'side_swing_angle': 45,
                'stable_time': 10,
                'angle_velocity': 5,
                'width_of_cloth': 50,
                'thickness_threshold': 800,
                'imaging_powers': 50,
                'sunlight_powers': 30,
                'eclipse_powers': 8,
                'maneuver_powers': 500,
                'downlink_rate': 4
            },
            'infrared': {
                'max_storage': 500,
                'max_battery': 5000,
                'resolution': 0.5,
                'pitch_angle': 45,
                'side_swing_angle': 45,
                'stable_time': 10,
                'angle_velocity': 5,
                'width_of_cloth': 50,
                'thickness_threshold': 800,
                'imaging_powers': 50,
                'sunlight_powers': 30,
                'eclipse_powers': 8,
                'maneuver_powers': 500,
                'downlink_rate': 4
            }
        }

    def generate_tasks(self, file_path):
        """
        从Excel表格批量生成用户任务
        :param file_path: Excel文件路径
        :return: 任务列表
        """
        print("---------------------接收前队列大小", self.task_queue.qsize(), "_----------------------")
        tasks = []

        if not file_path:
            print("没有选择文件")
            return tasks

        wb = None  # 供 finally 中关闭
        try:
            import openpyxl
            from datetime import datetime

            # 加载Excel文件
            wb = openpyxl.load_workbook(file_path)
            ws = wb.active  # 使用第一个sheet
            task = None
            # 从第二行开始读取任务（第一行是表头）
            for row in range(2, ws.max_row + 1):
                # 检查整行是否为空
                is_empty_row = True
                for col in range(1, ws.max_column + 1):
                    if ws.cell(row=row, column=col).value is not None:
                        is_empty_row = False
                        break
                if is_empty_row:
                    continue

                # 读取各个字段（单元格可能是数字，统一 str() 后再 strip，None 走默认值）
                is_emergency = (str(ws.cell(row=row, column=2).value).strip() == '是') \
                    if ws.cell(row=row, column=2).value is not None else False
                # 判断是否是紧急任务,如果是紧急任务，优先级设置为6
                priority = int(ws.cell(row=row, column=1).value) if ws.cell(row=row, column=1).value else 1
                priority = 6 if is_emergency else priority
                task_type = str(ws.cell(row=row, column=3).value).strip() if ws.cell(row=row,
                                                                                      column=3).value is not None else " "
                # sensor_type = ws.cell(row=row, column=4).value.strip() if ws.cell(row=row, column=4).value else (
                #     random.choice(['SAR', 'optical', 'infrared']))
                sensor_type = str(ws.cell(row=row, column=4).value).strip() if ws.cell(row=row,
                                                                                        column=4).value is not None else None
                resolution = float(ws.cell(row=row, column=5).value) if ws.cell(row=row, column=5).value else None
                cluster_name = str(ws.cell(row=row, column=8).value).strip() if ws.cell(row=row,
                                                                                         column=8).value is not None else None
                cycle_time = 60 * int(ws.cell(row=row, column=9).value) if ws.cell(row=row, column=9).value else 3600
                appoint_time = str(ws.cell(row=row, column=10).value).strip() if ws.cell(row=row,
                                                                                          column=10).value is not None else None
                appoint_time = _local_naive_to_utc(datetime.strptime(appoint_time, '%Y,%m,%d,%H,%M,%S')) if appoint_time else None  # 表格传入为本地时间，入口统一转 UTC
                cloud_thickness = float(ws.cell(row=row, column=11).value) if ws.cell(row=row, column=11).value else 0.0

                task_name = task_type + str(uuid.uuid4())[:6]

                # 解析时间范围
                time_range = str(ws.cell(row=row, column=6).value).strip() if ws.cell(row=row,
                                                                                       column=6).value is not None else None
                if time_range:
                    time_range = time_range.split('-')
                    start_time = _local_naive_to_utc(datetime.strptime(time_range[0].strip('()'), '%Y,%m,%d,%H,%M,%S'))  # 表格传入为本地时间，入口统一转 UTC
                    end_time = _local_naive_to_utc(datetime.strptime(time_range[1].strip('()'), '%Y,%m,%d,%H,%M,%S'))  # 表格传入为本地时间，入口统一转 UTC
                    start_time1 = start_time.replace(tzinfo=timezone.utc)
                    end_time1 = end_time.replace(tzinfo=timezone.utc)
                else:  # 如果没有指定时间范围，则默认使用整个仿真时间段
                    start_time = self.start_time
                    end_time = self.end_time
                    start_time1 = self.start_time.replace(tzinfo=timezone.utc)
                    end_time1 = self.end_time.replace(tzinfo=timezone.utc)
                # 解析区域坐标
                area_value = ws.cell(row=row, column=7).value
                if area_value is None:
                    print(f"Warning: Empty area value at row {row}, column 7")
                    continue
                area = str(area_value).strip()
                if task_type in ['区域目标', '广域目标', '陆地区域目标案例', '海洋搜救案例']:
                    areas = area.split('|')
                    boundary = []
                    for str_area in areas:
                        if str_area == '':
                            continue
                        locations = str_area.strip('[]').split(',')
                        boundary.append((float(locations[0].strip()), float(locations[1].strip())))

                    # 创建任务对象
                    task = Task(task_id=None, task_name=task_name, is_emergency=is_emergency, task_type=task_type,
                                sensor_type=sensor_type, resolution=resolution,
                                earliest_start_time=start_time1, latest_end_time=end_time1,
                                target_location=boundary, boundary_points=boundary,
                                cluster_name=cluster_name if cluster_name else None,
                                appointment_time=appoint_time, cloud_thickness=cloud_thickness)
                    # 将任务加入到数据库中
                    task.task_id = self._add_task(task, start_time, end_time, task_name, cloud_thickness)
                    # 如果任务为定时任务，则加入到定时任务列表中
                    if task.appointment_time:
                        print("读取到的预约时间：", task.appointment_time)
                        self.appointment_tasks.append(task)
                        continue
                    # 任务缓冲区中加入任务
                    self.tasks_buffer[task.task_id] = task
                    # 判断任务类型是否为区域目标或广域目标，如果是，则需要将任务拆分成多个点目标任务
                    subtasks = split_area_target_into_point_targets(task)
                    tasks.extend(subtasks)
                else:
                    area = area.strip('[]').split(',')
                    latitude = round(float(area[0].strip()), 2)
                    longitude = round(float(area[1].strip()), 2)
                    target_location = [latitude, longitude]
                    # 创建任务对象
                    task = Task(task_id=None, task_name=task_name, priority=priority, is_emergency=is_emergency,
                                task_type=task_type,
                                cycle_time=cycle_time if cycle_time else 3600,
                                sensor_type=sensor_type, resolution=resolution,
                                earliest_start_time=start_time1, latest_end_time=end_time1,
                                target_location=target_location, cluster_name=cluster_name if cluster_name else None,
                                appointment_time=appoint_time, cloud_thickness=cloud_thickness)
                    # 将任务加入到数据库中
                    task.task_id = self._add_task(task, start_time, end_time, task_name, cloud_thickness)
                    # 如果任务时定时任务，则加入到定时任务列表中
                    if task.appointment_time:
                        print("读取到的预约时间：", task.appointment_time)
                        self.appointment_tasks.append(task)
                        continue
                    if task_type == '周期观测':
                        # 任务缓冲区中加入任务
                        self.tasks_buffer[task.task_id] = task
                        subtasks = split_cycle_target_into_point_targets(task)
                        tasks.extend(subtasks)
                    else:
                        tasks.append(task)

            # with self.queue_lock:
            print("目前所有的主任务", self.tasks_buffer.keys())
            self.task_queue.put(tasks)
            print("---------------------接收后队列大小", self.task_queue.qsize(), "-----------------------")
            if tasks:
                return tasks
            return self.appointment_tasks

        except Exception as e:
            print(f"读取任务表格时发生错误：{str(e)}")
            import traceback
            print(traceback.format_exc())
            return []
        finally:
            # 无论成功或异常都关闭工作簿并删除临时Excel文件
            if wb is not None:
                wb.close()
            if file_path and os.path.exists(file_path):
                os.remove(file_path)

    def generate_single_task(self, data):
        """
        从用户处接收单个任务
        :param data: 任务数据
        :return: 任务对象列表
        """
        # task_id = data['task_id']
        priority = data['priority'] if data['priority'] else 1
        is_emergency = True if data['is_urgent'].strip() == "是" else False
        priority = 6 if is_emergency else priority
        task_type = data['type'].strip() if data['type'] else ""
        task_name = task_type + str(uuid.uuid4())[:6]
        sensor_type = data['payload'].strip() if data['payload'] else None
        resolution = float(data['resolution']) if data['resolution'] else None
        start_time = data['timeRanges'][0].strip().split(',')[0] if data['timeRanges'][0] else None
        end_time = data['timeRanges'][0].strip().split(',')[1] if data['timeRanges'][0] else None
        if start_time is None or end_time is None or start_time == "" or end_time == "":
            start_time = self.now_time
            end_time = start_time + timedelta(days=1)
        else:
            start_time = _local_naive_to_utc(datetime.strptime(start_time, '%Y-%m-%d %H:%M:%S'))  # 前端传入为本地时间，入口统一转 UTC
            end_time = _local_naive_to_utc(datetime.strptime(end_time, '%Y-%m-%d %H:%M:%S'))  # 前端传入为本地时间，入口统一转 UTC
        cycle_time = 60 * int(data['cycle'].strip()) if data['cycle'].strip() else 3600
        cloud_thickness = float(data['cloud_thickness']) if data['cloud_thickness'] else 0.0
        cluster_name = data['cluster_name'].strip() if data['cluster_name'] else None

        if data['appoint_time'] == "" or data['appoint_time'] is None:
            appoint_time = None
        else:
            appoint_time = _local_naive_to_utc(datetime.strptime(data['appoint_time'].strip(), '%Y-%m-%d %H:%M:%S')) if not (  # 前端传入为本地时间，入口统一转 UTC
                    data['appoint_time'] == "") else None
            print("读取到的预约时间：", appoint_time)

        target_location = data['coordinates']
        locations = []
        for target in target_location:
            area = target.strip().split(",")
            latitude = round(float(area[0].strip()), 2)
            longitude = round(float(area[1].strip()), 2)
            locations.append((latitude, longitude))

        boundary = locations
        if len(locations) == 1:
            locations = locations[0]
            boundary = None
        # 创建任务对象，并将其加入到待处理任务列表中（包装成列表，统一格式）
        tasks = []
        # 创建任务对象
        task = Task(task_id=None, task_name=task_name, priority=priority, is_emergency=is_emergency,
                    task_type=task_type,
                    sensor_type=sensor_type, cycle_time=cycle_time,
                    resolution=resolution, earliest_start_time=start_time.replace(tzinfo=timezone.utc),
                    latest_end_time=end_time.replace(tzinfo=timezone.utc), appointment_time=appoint_time,
                    target_location=locations, boundary_points=boundary,
                    cluster_name=cluster_name, cloud_thickness=cloud_thickness)

        # 将任务加入到数据库中
        task.task_id = self._add_task(task, start_time, end_time, task_name, cloud_thickness)

        if appoint_time:
            print("读取到的预约时间：", appoint_time)
            self.appointment_tasks.append(task)
            return
        if task_type in ['区域目标', '广域目标', '陆地区域目标案例', '海洋搜救案例']:
            # 任务缓冲区中加入任务
            self.tasks_buffer[task.task_id] = task
            # 判断任务类型是否为区域目标或广域目标，如果是，则需要将任务拆分成多个点目标任务
            subtasks = split_area_target_into_point_targets(task)
            tasks.extend(subtasks)
        else:
            if task_type == '周期观测':
                # 任务缓冲区中加入任务
                self.tasks_buffer[task.task_id] = task
                subtasks = split_cycle_target_into_point_targets(task)
                tasks.extend(subtasks)
            else:
                tasks.append(task)
        self.task_queue.put(tasks)
        print("目前所有的主任务", self.tasks_buffer.keys())
        return task

    def generate_tasks_by_page(self, data_list):
        """
        从用户处接收批量任务
        :param data_list: 任务数据
        :return: 任务对象列表
        """
        tasks = []
        for data in data_list:
            priority = data['priority']
            is_emergency = True if data['is_urgent'].strip() == "是" else False
            task_type = data['type'].strip()
            sensor_type = data['payload'].strip()
            resolution = float(data['resolution'])
            start_time = data['timeRanges'][0].strip().split(',')[0]
            end_time = data['timeRanges'][0].strip().split(',')[1]
            start_time = _local_naive_to_utc(datetime.strptime(start_time, '%Y-%m-%d %H:%M:%S'))  # 前端传入为本地时间，入口统一转 UTC
            end_time = _local_naive_to_utc(datetime.strptime(end_time, '%Y-%m-%d %H:%M:%S'))  # 前端传入为本地时间，入口统一转 UTC
            cycle_time = 60 * int(data['cycle'].strip()) if data['cycle'].strip() else 3600
            cloud_thickness = float(data['cloud_thickness']) if data['cloud_thickness'] else 0.0

            if data['appoint_time'] == "" or data['appoint_time'] is None:
                appoint_time = None
            else:
                appoint_time = _local_naive_to_utc(datetime.strptime(data['appoint_time'].strip(), '%Y-%m-%d %H:%M:%S')) if not (  # 前端传入为本地时间，入口统一转 UTC
                        data['appoint_time'] == "") else None
                print("读取到的预约时间：", appoint_time)

            target_location = data['coordinates']
            task_name = task_type + str(uuid.uuid4())[:6]
            locations = []
            for target in target_location:
                area = target.strip().split(",")
                latitude = round(float(area[0].strip()), 2)
                longitude = round(float(area[1].strip()), 2)
                locations.append([latitude, longitude])
            if len(locations) == 1:
                locations = locations[0]
            cluster_name = data['cluster_name'].strip()
            # 创建任务对象，并将其加入到待处理任务列表中（包装成列表，统一格式）
            # 创建任务对象
            task = Task(task_id=None, priority=priority, is_emergency=is_emergency, task_type=task_type,
                        sensor_type=sensor_type, cycle_time=cycle_time,
                        resolution=resolution, earliest_start_time=start_time.replace(tzinfo=timezone.utc),
                        latest_end_time=end_time.replace(tzinfo=timezone.utc), appointment_time=appoint_time,
                        target_location=locations, boundary_points=locations,
                        cluster_name=cluster_name, cloud_thickness=cloud_thickness)
            # 将任务加入到数据库中
            task.task_id = self._add_task(task, start_time, end_time, task_name, cloud_thickness)
            if appoint_time:
                print("读取到的预约时间：", appoint_time)
                self.appointment_tasks.append(task)
                return
            if task_type in ['区域目标', '广域目标', '陆地区域目标案例', '海洋搜救案例']:
                # 任务缓冲区中加入任务
                self.tasks_buffer[task.task_id] = task
                # 判断任务类型是否为区域目标或广域目标，如果是，则需要将任务拆分成多个点目标任务
                subtasks = split_area_target_into_point_targets(task)
                tasks.extend(subtasks)
            else:
                if task_type == '周期观测':
                    # 任务缓冲区中加入任务
                    self.tasks_buffer[task.task_id] = task
                    subtasks = split_cycle_target_into_point_targets(task)
                    tasks.extend(subtasks)
                else:
                    tasks.append(task)

        self.task_queue.put(tasks)
        print("目前所有的主任务", self.tasks_buffer.keys())
        return tasks

    def generate_random_tasks(self, point_task_count, area_task_count, moving_task_count, cluster_name, start_time,
                              end_time):
        """
        生成指定数量随机任务
        :param start_time: 开始时间
        :param end_time: 结束时间
        :param cluster_name: 星簇名
        :param point_task_count: 点目标任务数量
        :param area_task_count: 区域目标任务数量
        :param moving_task_count: 移动目标任务数量
        :return: 任务列表
        """
        print("数量：", point_task_count, area_task_count, moving_task_count)
        tasks = []
        # 生成指定数量的点目标任务
        parent_task_id = None
        cycle_time = None
        # 计算时间差
        time_diff = (end_time - start_time).total_seconds()
        if time_diff <= 0:
            return
        # 随机生成开始时间
        start_time = start_time + timedelta(seconds=random.randint(0, int(time_diff // 2)))
        end_time = end_time
        # 生成指定数量的点目标任务
        for i in range(point_task_count):
            priority = random.randint(1, 5)
            is_emergency = False
            sensor_type = random.choice(['optical', 'SAR', 'infrared'])
            task_type = "点目标"
            task_name = task_type + str(uuid.uuid4())[:6]
            resolution = random.choice([0.5, 1])
            start_time = start_time + timedelta(seconds=random.randint(0, int(time_diff // 2)))
            end_time = end_time
            target_location = [round(random.uniform(-90, 90), 2), round(random.uniform(-180, 180), 2)]
            cloud_thickness = random.randint(0, 1000)
            task = Task(task_id=None, priority=priority, is_emergency=is_emergency, task_type=task_type,
                        sensor_type=sensor_type, resolution=resolution,
                        earliest_start_time=start_time.replace(tzinfo=timezone.utc),
                        latest_end_time=end_time.replace(tzinfo=timezone.utc), target_location=target_location,
                        parent_task_id=parent_task_id, cycle_time=cycle_time, cluster_name=cluster_name,
                        cloud_thickness=cloud_thickness)
            task.task_id = self._add_task(task, start_time, end_time, task_name, cloud_thickness)
            tasks.append(task)
        # 生成指定数量的区域目标任务
        for i in range(area_task_count):
            priority = random.randint(1, 5)
            is_emergency = False
            sensor_type = random.choice(['optical', 'SAR', 'infrared'])
            task_type = "区域目标"
            task_name = task_type + str(uuid.uuid4())[:6]
            resolution = random.choice([0.5, 1])
            start_time = start_time + timedelta(seconds=random.randint(0, int(time_diff // 2)))
            end_time = end_time

            cloud_thickness = random.randint(0, 1000)

            # 生成随机的边界点集合
            boundary_points = []
            num_points = random.randint(3, 10)

            # 生成中心点
            center_lat = round(random.uniform(-80, 80), 2)  # 避免太靠近极地
            center_lon = round(random.uniform(-170, 170), 2)  # 避免跨越日期变更线

            # 生成围绕中心点的其他点
            for j in range(num_points):
                # 计算当前点的角度（均匀分布在圆周上）
                angle = 2 * math.pi * j / num_points
                # 随机半径（确保多边形不会太大或太小）
                radius = random.uniform(0.5, 5.0)
                # 计算相对于中心点的偏移
                lat_offset = radius * math.cos(angle)
                lon_offset = radius * math.sin(angle)
                # 计算实际经纬度
                latitude = round(center_lat + lat_offset, 2)
                longitude = round(center_lon + lon_offset, 2)
                # 确保经纬度在有效范围内
                latitude = max(-90, min(90, latitude))
                longitude = max(-180, min(180, longitude))
                boundary_points.append([latitude, longitude])

            # # 添加第一个点作为最后一个点，确保完全封闭
            # boundary_points.append(boundary_points[0])
            # 创建任务对象
            task = Task(
                task_id=None, priority=priority, is_emergency=is_emergency, task_type=task_type,
                sensor_type=sensor_type, boundary_points=boundary_points,
                resolution=resolution, earliest_start_time=start_time.replace(tzinfo=timezone.utc),
                latest_end_time=end_time.replace(tzinfo=timezone.utc),
                target_location=boundary_points, cycle_time=cycle_time, cluster_name=cluster_name,
                cloud_thickness=cloud_thickness
            )
            # 将任务加入到数据库中
            task.task_id = self._add_task(task, start_time, end_time, task_name, cloud_thickness)
            # 任务缓冲区中加入任务
            self.tasks_buffer[task.task_id] = task
            # 需要将任务拆分成多个点目标任务
            subtasks = split_area_target_into_point_targets(task)
            tasks.extend(subtasks)
        # 生成指定数量的移动目标任务
        for i in range(moving_task_count):
            priority = random.randint(1, 5)
            is_emergency = False
            sensor_type = random.choice(['optical', 'SAR', 'infrared'])
            task_type = "移动目标"
            task_name = task_type + str(uuid.uuid4())[:6]
            resolution = random.choice([0.5, 1])
            start_time = start_time + timedelta(seconds=random.randint(0, int(time_diff // 2)))
            end_time = end_time
            cloud_thickness = random.randint(0, 1000)

            target_location = [round(random.uniform(-90, 90), 2), round(random.uniform(-180, 180), 2)]
            task = Task(
                task_id=None, priority=priority,
                is_emergency=is_emergency,
                task_type=task_type,
                sensor_type=sensor_type, resolution=resolution,
                earliest_start_time=start_time.replace(tzinfo=timezone.utc),
                latest_end_time=end_time.replace(tzinfo=timezone.utc),
                target_location=target_location,
                parent_task_id=parent_task_id, cycle_time=cycle_time,
                cluster_name=cluster_name,
                cloud_thickness=cloud_thickness
            )
            task.task_id = self._add_task(task, start_time, end_time, task_name, cloud_thickness)
            tasks.append(task)
        self.task_queue.put(tasks)

    def _add_task(self, task, start_time, end_time, task_name, cloud_thickness):
        """
            将任务添加到数据库中
        """
        with self.app.app_context():
            # 转换skyfield.Time对象为datetime
            def convert_skyfield_time(time_obj):
                if hasattr(time_obj, 'utc_datetime'):
                    return time_obj.utc_datetime()
                return None

            task_model = NewTaskModel(
                id=task.task_id,
                priority=task.priority,
                task_name=task_name,
                is_emergency=task.is_emergency,
                sensor_type=task.sensor_type,
                task_type=task.task_type,
                resolution=task.resolution,
                start_time=start_time,
                end_time=end_time,
                user_start_time=start_time,
                user_end_time=end_time,
                appoint_time=task.appointment_time,
                target_location=str(task.target_location),
                assigned_satellite_name=None,
                status="等待规划",
                cluster_name=task.cluster_name,
                friend_task_id=None,
                cloud_thickness=cloud_thickness
            )

            db.session.add(task_model)
            db.session.commit()
            return task_model.id

    def delete_task(self, task_id, status):
        # 加锁保护共享任务列表/字典（与网络线程 check_execute_tasks 共用同一把 task_lock）
        with self.task_lock:
            if status == "等待规划":
                for task in self.appointment_tasks:
                    if task.task_id == task_id:
                        self.appointment_tasks.remove(task)
                        break
            if status == "等待执行":
                if task_id in self.satellite_network.net_tasks_buffer:
                    main_task = self.satellite_network.net_tasks_buffer[task_id]
                    for task in main_task.subtasks:
                        self.satellite_network.new_tasks.remove(task)
                        if task.assigned_satellite:
                            self.satellite_network.satellites[task.assigned_satellite].tasks_len -= 1
                    del self.satellite_network.net_tasks_buffer[task_id]
                else:
                    for task in self.satellite_network.new_tasks:
                        if task.task_id == task_id:
                            self.satellite_network.new_tasks.remove(task)
                            if task.assigned_satellite:
                                self.satellite_network.satellites[task.assigned_satellite].tasks_len -= 1
                            break
        # 数据库操作不放在 task_lock 内（SQLAlchemy scoped session 自带线程语义）
        # 与 manual_end_task 归档逻辑保持一致：删除时归档到旧任务表（状态标记 Failed，保留原结束时间）
        with self.app.app_context():
            new_task = NewTaskModel.query.filter_by(id=task_id).first()
            if new_task:
                old_task = OldTaskModel(id=new_task.id, priority=new_task.priority,
                                        task_name=new_task.task_name,
                                        is_emergency=new_task.is_emergency,
                                        task_type=new_task.task_type, sensor_type=new_task.sensor_type,
                                        resolution=new_task.resolution, start_time=new_task.start_time,
                                        end_time=new_task.end_time, appoint_time=new_task.appoint_time,
                                        target_location=new_task.target_location,
                                        assigned_satellite_name=new_task.assigned_satellite_name,
                                        status="Failed", is_photo=True,
                                        path=None, cloud_thickness=new_task.cloud_thickness)
                db.session.add(old_task)
                db.session.delete(new_task)
                db.session.commit()

    def edit_task(self, task_id, data):
        """
        编辑未完成任务：同步更新内存中的任务对象并写入数据库
        :param task_id: 任务id
        :param data: 需要更新的字段（与 generate_single_task 入参格式一致，只更新传入的字段）
        """
        # 解析需要更新的字段（解析逻辑与 generate_single_task 保持一致）
        updates = {}
        if data.get('task_name'):
            updates['task_name'] = data['task_name']
        if data.get('is_urgent') is not None and data['is_urgent'].strip():
            updates['is_emergency'] = True if data['is_urgent'].strip() == "是" else False
        if data.get('priority'):
            updates['priority'] = int(data['priority'])
        if updates.get('is_emergency'):
            updates['priority'] = 6  # 与 generate_single_task 一致：紧急任务优先级为6
        if data.get('type') is not None and data['type'].strip():
            updates['task_type'] = data['type'].strip()
        if data.get('payload') is not None and data['payload'].strip():
            updates['sensor_type'] = data['payload'].strip()
        if data.get('resolution'):
            updates['resolution'] = float(data['resolution'])
        if data.get('timeRanges') and data['timeRanges'][0]:
            start_time = data['timeRanges'][0].strip().split(',')[0]
            end_time = data['timeRanges'][0].strip().split(',')[1]
            if start_time and end_time:
                # 前端传入为本地时间，入口统一转 UTC
                updates['start_time'] = _local_naive_to_utc(datetime.strptime(start_time, '%Y-%m-%d %H:%M:%S'))
                updates['end_time'] = _local_naive_to_utc(datetime.strptime(end_time, '%Y-%m-%d %H:%M:%S'))
        if 'appoint_time' in data:
            if data['appoint_time'] == "" or data['appoint_time'] is None:
                updates['appoint_time'] = None
            else:
                updates['appoint_time'] = _local_naive_to_utc(
                    datetime.strptime(data['appoint_time'].strip(), '%Y-%m-%d %H:%M:%S'))  # 前端传入为本地时间，入口统一转 UTC
        if data.get('cycle') and data['cycle'].strip():
            updates['cycle_time'] = 60 * int(data['cycle'].strip())
        if data.get('cloud_thickness'):
            updates['cloud_thickness'] = float(data['cloud_thickness'])
        if 'cluster_name' in data:
            updates['cluster_name'] = data['cluster_name'].strip() if data['cluster_name'] else None
        if data.get('coordinates'):
            locations = []
            for target in data['coordinates']:
                area = target.strip().split(",")
                latitude = round(float(area[0].strip()), 2)
                longitude = round(float(area[1].strip()), 2)
                locations.append((latitude, longitude))
            boundary = locations
            if len(locations) == 1:
                locations = locations[0]
                boundary = None
            updates['target_location'] = locations
            updates['boundary_points'] = boundary

        # 加锁保护共享任务列表/字典（与 delete_task 共用同一把 task_lock），同步更新内存中的任务对象
        with self.task_lock:
            for task in self._find_memory_tasks(task_id):
                for key, value in updates.items():
                    # 数据库字段名与内存任务对象属性名的映射
                    attr = {'start_time': 'earliest_start_time', 'end_time': 'latest_end_time',
                            'appoint_time': 'appointment_time'}.get(key, key)
                    if attr in ('earliest_start_time', 'latest_end_time') and value is not None:
                        value = value.replace(tzinfo=timezone.utc)
                        # 同步时间备份字段，供重规划时恢复使用
                        setattr(task, attr + '_backup', value)
                    setattr(task, attr, value)

        # 数据库操作不放在 task_lock 内（与 delete_task 一致）
        with self.app.app_context():
            new_task = NewTaskModel.query.filter_by(id=task_id).first()
            if not new_task:
                return
            for key, value in updates.items():
                if key == 'boundary_points' or key == 'cycle_time':
                    continue  # 内存对象专有字段，数据库无对应列
                if key == 'target_location':
                    new_task.target_location = str(value)  # 与 _add_task 一致，存储 str(task.target_location)
                elif key == 'start_time':
                    new_task.start_time = value
                    new_task.user_start_time = value
                elif key == 'end_time':
                    new_task.end_time = value
                    new_task.user_end_time = value
                else:
                    setattr(new_task, key, value)
            db.session.commit()

    def _find_memory_tasks(self, task_id):
        """
        查找内存中持有的指定任务对象（预约任务列表、任务缓冲区、网络任务缓冲区、新任务/暂停任务列表）
        :param task_id: 任务id
        :return: 任务对象列表
        """
        tasks = []
        for task in self.appointment_tasks:
            if task.task_id == task_id:
                tasks.append(task)
        if task_id in self.tasks_buffer:
            tasks.append(self.tasks_buffer[task_id])
        if self.satellite_network:
            if task_id in self.satellite_network.net_tasks_buffer:
                tasks.append(self.satellite_network.net_tasks_buffer[task_id])
            for task in self.satellite_network.new_tasks:
                if task.task_id == task_id:
                    tasks.append(task)
            for task in self.satellite_network.pause_tasks:
                if task.task_id == task_id:
                    tasks.append(task)
        return tasks

    def manual_end_task(self, task_id, status):
        # 加锁保护共享任务列表/字典（与网络线程共用同一把 task_lock）
        with self.task_lock:
            if status == "等待规划":
                for task in self.appointment_tasks:
                    if task.task_id == task_id:
                        self.appointment_tasks.remove(task)
                        break
            if status == "等待执行":
                if task_id in self.satellite_network.net_tasks_buffer:
                    main_task = self.satellite_network.net_tasks_buffer[task_id]
                    for task in main_task.subtasks:
                        self.satellite_network.new_tasks.remove(task)
                        if task.assigned_satellite:
                            self.satellite_network.satellites[task.assigned_satellite].tasks_len -= 1
                    del self.satellite_network.net_tasks_buffer[task_id]
                else:
                    for task in self.satellite_network.new_tasks:
                        if task.task_id == task_id:
                            self.satellite_network.new_tasks.remove(task)
                            if task.assigned_satellite:
                                self.satellite_network.satellites[task.assigned_satellite].tasks_len -= 1
                            break
            if status == "暂停":
                if task_id in self.satellite_network.net_tasks_buffer:
                    main_task = self.satellite_network.net_tasks_buffer[task_id]
                    for task in main_task.subtasks:
                        self.satellite_network.pause_tasks.remove(task)
                        if task.assigned_satellite:
                            self.satellite_network.satellites[task.assigned_satellite].tasks_len -= 1
                    del self.satellite_network.net_tasks_buffer[task_id]
                else:
                    for task in self.satellite_network.pause_tasks:
                        if task.task_id == task_id:
                            self.satellite_network.pause_tasks.remove(task)
                            if task.assigned_satellite:
                                self.satellite_network.satellites[task.assigned_satellite].tasks_len -= 1
                            break
        # 数据库操作不放在 task_lock 内
        with self.app.app_context():
            new_task = NewTaskModel.query.filter_by(id=task_id).first()
            if new_task:
                old_task = OldTaskModel(id=new_task.id, priority=new_task.priority,
                                        task_name=new_task.task_name,
                                        is_emergency=new_task.is_emergency,
                                        task_type=new_task.task_type, sensor_type=new_task.sensor_type,
                                        resolution=new_task.resolution, start_time=new_task.start_time,
                                        end_time=new_task.end_time, appoint_time=new_task.appoint_time,
                                        target_location=new_task.target_location,
                                        assigned_satellite_name=new_task.assigned_satellite_name,
                                        status="Failed", is_photo=True,
                                        path=None, cloud_thickness=new_task.cloud_thickness)
                db.session.add(old_task)
                db.session.delete(new_task)
                db.session.commit()

    # 任务管理
    def receive_tasks(self):
        """
        从用户处接收任务
        """
        # 获取任务队列中的任务
        tasks = []
        al_tasks = []
        if not self.task_queue.empty():
            tasks = self.task_queue.get()
            print(f"运控中心接收到{len(tasks)}个任务: ")
        # 遍历预约任务列表，如果预约任务的开始时间与当前时间相差 5秒*时间倍速 以内，则将其加入任务列表中
        with self.task_lock:
            appointment_snapshot = list(self.appointment_tasks)  # 锁内拍快照，避免遍历时被Flask线程修改
        for task in appointment_snapshot:
            # print("任务的预约时间", task.appointment_time.replace(tzinfo=None))
            if (task.appointment_time.replace(
                    tzinfo=None) - self.now_time).total_seconds() <= INTER_VAL_TIME * self.time_multiple:
                al_tasks.append(task)
                if task.task_type in ['区域目标', '广域目标', '陆地区域目标案例', '海洋搜救案例']:
                    # 任务缓冲区中加入任务
                    self.tasks_buffer[task.task_id] = task
                    # 判断任务类型是否为区域目标或广域目标，如果是，则需要将任务拆分成多个点目标任务
                    subtasks = split_area_target_into_point_targets(task)
                    tasks.extend(subtasks)
                elif task.task_type == '周期观测':
                    # 任务缓冲区中加入任务
                    self.tasks_buffer[task.task_id] = task
                    subtasks = split_cycle_target_into_point_targets(task)
                    tasks.extend(subtasks)
                else:
                    tasks.append(task)
        with self.task_lock:
            for task in al_tasks:
                if task in self.appointment_tasks:  # 可能已被delete_task/manual_end_task并发移除
                    self.appointment_tasks.remove(task)
        return tasks

    # 卫星网络状态管理
    def get_net_state(self):
        """
        获取卫星网络状态
        """
        if self.satellite_network:
            self.network_state = self.satellite_network.network_state
            return self.network_state
        return None

    # 规划任务
    def planning_tasks(self, tasks, model):
        """
        规划任务（简单实现）
        输入：任务，卫星网络状态
        输出：任务部署方案
        """
        count = 0
        original_count = 0
        main_set = set()
        self.satellite_network.is_planed = False  # 设置网络状态为规划中
        self.time_multiple = 1  # 规划前设置为1
        try:

            print("收到的原始任务数目：", len(tasks))

            with ((self.app.app_context())):
                # 合并任务
                i = 0
                processed_tasks = []
                while i < len(tasks):
                    task1 = tasks[i]
                    #   如果是重规划任务，则直接加入
                    if task1.is_replan:
                        processed_tasks.append(task1)
                        i += 1
                        continue
                    if task1.parent_task_id or task1.friend_task:
                        if task1.parent_task_id:
                            processed_tasks.append(task1)
                            main_task = self.tasks_buffer.get(task1.parent_task_id)
                            if main_task is not None:  # 父任务可能已被移除，避免 KeyError
                                main_set.add(main_task)
                        i += 1
                        continue
                    else:
                        j = i + 1
                        while j < len(tasks):
                            task2 = tasks[j]
                            if not task2.parent_task_id and not task2.friend_task:
                                # 判断合并条件
                                if (task1.sensor_type == task2.sensor_type and
                                        task1.earliest_start_time == task2.earliest_start_time and
                                        task1.latest_end_time == task2.latest_end_time and
                                        ((task1.target_location[0] - task2.target_location[0]) ** 2 +
                                         (task1.target_location[1] - task2.target_location[1]) ** 2 <= 8)):
                                    # 合并任务，双向映射，并把第二者删掉
                                    task1.friend_task = task2.task_id
                                    task2.friend_task = task1.task_id

                                    new_task_model = NewTaskModel.query.filter_by(id=task1.task_id).first()
                                    if new_task_model:
                                        new_task_model.friend_task_id = task1.friend_task
                                        db.session.commit()

                                    # 第二个任务从数据库删除
                                    newtask_model = NewTaskModel.query.filter_by(id=task2.task_id).first()
                                    if newtask_model:
                                        oldtask_model = OldTaskModel(
                                            id=newtask_model.id,
                                            task_name=newtask_model.task_name,
                                            priority=newtask_model.priority,
                                            is_emergency=newtask_model.is_emergency,
                                            sensor_type=newtask_model.sensor_type,
                                            task_type=newtask_model.task_type,
                                            resolution=newtask_model.resolution,
                                            start_time=newtask_model.start_time,
                                            end_time=newtask_model.end_time,
                                            appoint_time=newtask_model.appoint_time,
                                            target_location=newtask_model.target_location,
                                            assigned_satellite_name=newtask_model.assigned_satellite_name,
                                            status="合并",
                                            friend_task_id=task1.task_id,
                                            is_photo=True,
                                            cloud_thickness=task1.cloud_thickness
                                        )
                                        db.session.delete(newtask_model)
                                        db.session.add(oldtask_model)
                                        db.session.commit()

                                    break
                            j += 1

                        processed_tasks.append(task1)
                        i += 1

                tasks = processed_tasks  # 最终规划的是合并后的任务
                print("预处理后的任务数目：", len(tasks))
                if not tasks:
                    print("传进来的任务是空的！！！！！")

                # 统计规划开始时间（真实墙钟，仅用于与 planing_end 求差得到规划耗时，两者同为本地时间，差值与时区无关，故不做 UTC 转换）
                planing_start = datetime.now()

                print("正在进行时间窗口计算")
                # 遍历任务，为每个任务分配卫星
                for task in tasks:
                    # 为任务计算时间窗口
                    # 判断任务是否有传感器类型，如果有，则只为该类型的卫星分配时间窗口
                    if task.sensor_type:
                        for sat_name, sat in self.satellite_network.satellites.items():
                            # 检查卫星是否满足任务要求
                            # 判断卫星是否属于任务指定的星簇（如任务没有指定星簇则忽略此条件）
                            # 并且卫星搭载的传感器类型与任务要求的传感器类型匹配
                            if ((not task.cluster_name or task.cluster_name in sat.belong_cluster) and
                                    task.sensor_type == sat.star_payload and task.resolution == sat.resolution_capability):
                                # print("根据载荷筛选", task.sensor_type, sat.star_payload)
                                sat.calculate_time_windows(task)
                    else:  # 如果任务没有传感器类型，则用所有卫星计算时间窗口
                        for sat_name, sat in self.satellite_network.satellites.items():
                            if not task.cluster_name or task.cluster_name in sat.belong_cluster:
                                # print("只根据星簇筛选", task.cluster_name, sat.belong_cluster)
                                sat.calculate_time_windows(task)
                        # sat.simple_window(task)
                    if task.parent_task_id is None:
                        original_count += 1
                    # 将任务的属性追加进一个json文件
                    # 收集所有任务的属性
                    # tasks_data = []
                    # for t in tasks:
                    #     def convert_datetime(obj):
                    #         if isinstance(obj, datetime):
                    #             return obj.isoformat()
                    #         elif isinstance(obj, list):
                    #             return [convert_datetime(item) for item in obj]
                    #         elif isinstance(obj, dict):
                    #             return {k: convert_datetime(v) for k, v in obj.items()}
                    #         return obj
                    #
                    #     task_dict = {
                    #         'task_id': getattr(t, 'task_id', None),
                    #         'priority': getattr(t, 'priority', None),
                    #         'is_emergency': getattr(t, 'is_emergency', None),
                    #         'task_type': getattr(t, 'task_type', None),
                    #         'sensor_type': getattr(t, 'sensor_type', None),
                    #         'resolution': getattr(t, 'resolution', None),
                    #         'earliest_start_time': convert_datetime(getattr(t, 'earliest_start_time', None)),
                    #         'latest_end_time': convert_datetime(getattr(t, 'latest_end_time', None)),
                    #         'appointment_time': convert_datetime(getattr(t, 'appointment_time', None)),
                    #         'target_location': getattr(t, 'target_location', None),
                    #         'boundary_points': getattr(t, 'boundary_points', None),
                    #         'cluster_name': getattr(t, 'cluster_name', None),
                    #         'cloud_thickness': getattr(t, 'cloud_thickness', None),
                    #         'parent_task_id': getattr(t, 'parent_task_id', None),
                    #         'friend_task': getattr(t, 'friend_task', None),
                    #         'assigned_satellite': getattr(t, 'assigned_satellite', None),
                    #         'status': getattr(t, 'status', None),
                    #         'visible_windows': convert_datetime(getattr(t, 'visible_windows', None)),
                    #     }
                    #     tasks_data.append(task_dict)
                    # # 追加写入json文件
                    # output_dir = os.path.join(os.getcwd(), 'output_plans')
                    # os.makedirs(output_dir, exist_ok=True)
                    # output_path = os.path.join(output_dir, 'tasks_timewindow.json')
                    # # 先读取原有内容
                    # old_data = []
                    # if os.path.exists(output_path):
                    #     try:
                    #         with open(output_path, 'r', encoding='utf-8') as f:
                    #             old_data = json.load(f)
                    #             if not isinstance(old_data, list):
                    #                 old_data = []
                    #     except Exception:
                    #         old_data = []
                    # # 合并并写回
                    # all_data = old_data + tasks_data
                    # with open(output_path, 'w', encoding='utf-8') as f:
                    #     json.dump(all_data, f, ensure_ascii=False, indent=2)
                    # 打印任务的可见时间窗口
                    # if task.appointment_time:
                    #     print(task.visible_windows)
                    print(task.task_id, task.visible_windows)

                print("时间窗口计算完成")

                start_statistics_time = self.now_time.replace(tzinfo=timezone.utc)
                end_statistics_time = self.now_time.replace(tzinfo=timezone.utc) + timedelta(days=1)
                diff_time = (end_statistics_time - start_statistics_time).total_seconds()

                # 调用遗传算法进行规划
                schedule = exportcfuc(tasks, list(self.satellite_network.satellites.values()),
                                      start_statistics_time, end_statistics_time,
                                      self.completed_gravity, self.balance_gravity,
                                      self.priority_gravity)

                # 统计规划结束时间（与 planing_start 配对求耗时，同为本地墙钟，不与仿真 UTC 时间比较，保持原样）
                planing_end = datetime.now()
                duration = (planing_end - planing_start).total_seconds()
                print("+++++++++++++++++规划耗时----------------------:", duration)
                # 最终要执行的任务
                new_tasks = []
                if schedule:
                    # print("统计数据")
                    # print("任务受理数目", schedule[model]["entry_tasks"])
                    # print("任务满足率", schedule[model]["task_satisfaction"])
                    # print("资源利用率", schedule[model]["resource_utilization"])
                    # print("成像质量", schedule[model]["imaging_quality"])
                    # print("存储消耗", schedule[model]["overall_storage_cost"])
                    # print("电量消耗", schedule[model]["overall_battery_cost"])
                    # print("遗传算法", schedule[4]["greedy"])
                    print("方案：", schedule[model]["schedule"])
                    if not schedule[model]["schedule"]:
                        print("没有可行方案!!!!!!!!!!!!!")
                    for item in schedule[model]["schedule"]:
                        is_found = False
                        # print(item["task_id"])
                        for task in tasks:
                            if task.task_id == item["task_id"]:
                                is_found = True
                                # print("找到匹配的任务：", "任务id：", task.task_id)
                                # 如果卫星的任务数目已达到上限，则跳过该任务
                                if self.satellite_network.satellites[item["satellite_id"]].tasks_len > 30:
                                    print(self.satellite_network.satellites[item["satellite_id"]],
                                          "卫星的任务数目已达到上限")
                                    break
                                self.satellite_network.satellites[item["satellite_id"]].tasks_len += 1
                                task.assigned_satellite = item["satellite_id"]
                                task.sensor_type = self.satellite_network.satellites[task.assigned_satellite].star_payload
                                task.resolution = self.satellite_network.satellites[
                                    task.assigned_satellite].resolution_capability
                                task.status = "等待执行"
                                task.earliest_start_time = item["start_time"]
                                task.latest_end_time = item["end_time"]
                                # task.battery_before_task = item["battery_before_task"]
                                task.battery_after_task = item["battery_after_task"]
                                # task.storge_before_task = item["storage_cap"]-item["storage_before_task"]
                                # task.storge_after_task = item["storage_cap"]-item["current_storage"]
                                task.current_storage = item["current_storage"]
                                task.execution_time = item["observation_time"] + item["maneuver_duration"]
                                task.target_attitude = item["attitude_after_task"]
                                # if item["light_power"] != 0:
                                #     task.light_power = item["light_power"]
                                ts = load.timescale()
                                # print("任务的起止时间：", task.earliest_start_time, task.latest_end_time)
                                t = ts.utc(task.earliest_start_time.replace(tzinfo=timezone.utc))
                                sat = self.satellite_network.satellites[task.assigned_satellite]
                                geocentric = sat.sat_model.at(t)
                                observatory = wgs84.subpoint(geocentric)
                                sub_point = [round(observatory.latitude.degrees, 2),
                                             round(observatory.longitude.degrees, 2)]

                                print("任务的id", task.task_id, "任务的经纬度：", task.target_location, "卫星的名字：",
                                      task.assigned_satellite,
                                      "卫星的经纬度：", sub_point)
                                # print(sat.tle_line1)
                                # print(sat.tle_line2)

                                # print("姿态：", item["attitude_after_task"])
                                # 如果任务是子任务，则将其从父任务的子任务id列表中删除(用于判断是否子任务是否全部可行)
                                if task.parent_task_id and task.parent_task_id in self.tasks_buffer and task.task_id in \
                                        self.tasks_buffer[task.parent_task_id].subtask_ids:
                                    self.tasks_buffer[task.parent_task_id].subtask_ids.remove(task.task_id)
                                else:  # 如果任务不是子任务，则将其加入新任务列表和数据库
                                    count += 1
                                    # 更新数据库
                                    if task.appointment_time:
                                        print("定时任务的不是子任务")
                                    newtask_model = NewTaskModel.query.filter_by(id=task.task_id).first()
                                    if newtask_model:
                                        # print("定时任务的数据库更新")
                                        newtask_model.status = task.status
                                        newtask_model.sensor_type = task.sensor_type
                                        newtask_model.resolution = task.resolution
                                        newtask_model.start_time = task.earliest_start_time
                                        newtask_model.end_time = task.latest_end_time
                                        newtask_model.assigned_satellite_name = task.assigned_satellite
                                        newtask_model.comment = f"{task.assigned_satellite}|{task.earliest_start_time}|{task.latest_end_time}"
                                        db.session.commit()
                                        new_tasks.append(task)
                                tasks.remove(task)
                                break

                        if not is_found:
                            # print("未找到匹配的任务！！！！！！")
                            for t in tasks:
                                print("剩余的任务：", t.task_id)
                print(self.tasks_buffer.keys())

                # 对任务缓冲区中的主任务进行处理，如果其子任务全部可行，则将其子任务加入新任务列表，否则将主任务加入失败任务列表
                main_tasks = list(main_set)  # 创建值的副本
                original_count += len(main_tasks)
                for main_task in main_tasks:
                    # print("主任务的的子任务长度：", main_task.task_id, len(main_task.subtask_ids))
                    if len(main_task.subtask_ids) == 0:
                        print("+++++++++++成功的主任务：", main_task.task_id)
                        count += 1  # 成功任务加1

                        new_tasks.extend(main_task.subtasks)  # 将子任务加入新任务列表

                        start_time = None
                        end_time = None
                        string = ""
                        comment = ""
                        for subtask in main_task.subtasks:
                            comment += f"{subtask.assigned_satellite}|{subtask.earliest_start_time}|{subtask.latest_end_time}" if comment == "" else f",{subtask.assigned_satellite}|{subtask.earliest_start_time}|{subtask.latest_end_time}"
                            # 更新任务的开始时间和结束时间,以所有子任务的最早的开始时间和最晚结束时间为准
                            if start_time is None or subtask.earliest_start_time < start_time:
                                start_time = subtask.earliest_start_time
                            if end_time is None or subtask.latest_end_time > end_time:
                                end_time = subtask.latest_end_time
                            string += f"{subtask.assigned_satellite}" if string == "" else f",{subtask.assigned_satellite}"
                            main_task.subtask_ids.append(subtask.task_id)
                        main_task.earliest_start_time = start_time
                        main_task.latest_end_time = end_time
                        main_task.comment = comment
                        main_task.assigned_satellite = string
                        main_task.status = "等待执行"
                        # 更新数据库
                        newtask_model = NewTaskModel.query.filter_by(id=main_task.task_id).first()
                        if newtask_model:
                            newtask_model.status = "等待执行"
                            newtask_model.start_time = main_task.earliest_start_time
                            newtask_model.end_time = main_task.latest_end_time
                            newtask_model.assigned_satellite_name = main_task.assigned_satellite
                            newtask_model.comment = main_task.comment
                            db.session.commit()
                            #  将规划成功的主任务加入网络任务缓冲区
                            self.satellite_network.net_tasks_buffer[main_task.task_id] = main_task
                    else:
                        tasks.append(main_task)  # 将主任务加入失败任务列表
                    self.tasks_buffer.pop(main_task.task_id)

                    # 将规划失败的任务从新任务数据库中删除，并添加到旧任务数据库中
                for task in tasks:
                    task.status = "Failed"
                    newtask_model = NewTaskModel.query.filter_by(id=task.task_id).first()
                    if newtask_model:
                        print("失败的任务名字：", newtask_model.task_name, newtask_model.assigned_satellite_name)
                        oldtask_model = OldTaskModel(
                            id=newtask_model.id,
                            task_name=newtask_model.task_name,
                            priority=newtask_model.priority,
                            is_emergency=newtask_model.is_emergency,
                            sensor_type=newtask_model.sensor_type,
                            task_type=newtask_model.task_type,
                            resolution=newtask_model.resolution,
                            start_time=newtask_model.start_time,
                            end_time=newtask_model.end_time,
                            appoint_time=newtask_model.appoint_time,
                            target_location=newtask_model.target_location,
                            assigned_satellite_name=newtask_model.assigned_satellite_name,
                            status=task.status,
                            is_photo=True,
                            cloud_thickness=newtask_model.cloud_thickness
                        )
                        db.session.delete(newtask_model)
                        db.session.add(oldtask_model)
                        db.session.commit()
                # 如果有规划结果，则直接添加统计数据
                if schedule:
                    schedule[model]["cluster_status"]["time"] = str(self.now_time)
                    for cluster in self.satellite_network.clusters:
                        if cluster.name not in schedule[model]["cluster_status"]:
                            schedule[model]["cluster_status"][cluster.name] = {
                                "cluster_time_cost": 0,
                                "cluster_battery_utilization": 0,
                                "cluster_storage_utilization": 0,
                                "cluster_battery_cost": 0,
                                "cluster_storage_cost": 0,
                                "time_occupancy_ratio": 0,
                                "cluster_tasks_count": 0
                            }
                        else:
                            # 计算星簇占用时间比例，使用星簇的占用时间/星簇的卫星数*时间差
                            denominator = diff_time * cluster.satellite_count
                            if denominator:  # 除零保护：分母为0时跳过该项
                                schedule[model]["cluster_status"][cluster.name]["time_occupancy_ratio"] = \
                                    schedule[model]["cluster_status"][cluster.name]["cluster_time_cost"] / denominator
                    data = {
                        "time": str(self.now_time),  # 当前时间
                        "duration": duration,  # 规划耗时
                        'overall_battery_cost': schedule[model]["overall_battery_cost"],  # 总电池消耗
                        'overall_storage_cost': schedule[model]["overall_storage_cost"],  # 总存储消耗
                        "task_satisfaction": schedule[model]["task_satisfaction"],  # 任务满足率
                        "genetic": schedule[4]["genetic"],  # 遗传算法
                        "greedy": schedule[4]["greedy"],  # 贪心算法
                        "ant_colony": schedule[4]["ant_colony"]  # 蚁群算法
                    }
                    #  添加统计数据
                    self.cluster_data.append(schedule[model]["cluster_status"])
                    self.statistical_data.append(data)
                    self.algorithm_name = schedule[model]["algorithm"]
                else:  # 如果没有规划结果，则添加默认统计数据
                    cluster_dict = {"time": str(self.now_time)}
                    for cluster in self.satellite_network.clusters:
                        cluster_dict[cluster.name] = {
                            "cluster_time_cost": 0,
                            "cluster_battery_utilization": 0,
                            "cluster_storage_utilization": 0,
                            "cluster_battery_cost": 0,
                            "cluster_storage_cost": 0,
                            "time_occupancy_ratio": 0,
                            "cluster_tasks_count": 0
                        }
                    data = {
                        "time": str(self.now_time),  # 当前时间
                        "duration": duration,  # 规划耗时
                        'overall_battery_cost': 0,  # 总电池消耗
                        'overall_storage_cost': 0,  # 总存储消耗
                        "task_satisfaction": 0,  # 任务满足率
                        "genetic": {'task_satisfaction': 0, 'overall_battery_cost': 0, 'overall_storage_cost': 0},  # 遗传算法
                        "greedy": {'task_satisfaction': 0, 'overall_battery_cost': 0, 'overall_storage_cost': 0},  # 贪心算法
                        "ant_colony": {'task_satisfaction': 0, 'overall_battery_cost': 0, 'overall_storage_cost': 0}  # 蚁群算法
                    }
                    self.cluster_data.append(cluster_dict)
                    self.statistical_data.append(data)
                    self.algorithm_name = None

                # for t in new_tasks:
                # print("任务：", t.task_id, t.parent_task_id)
                if schedule and schedule[-1]:
                    planned_data = schedule[-1]
                    print(planned_data)
                    planned_data = self.flatten_data(planned_data)
                    print(planned_data)
                    self.planning_results.extend(planned_data)
            return new_tasks
        except Exception as e:
            # 规划过程异常兜底：打印异常信息，避免 is_planed 一直为 False 导致仿真假死
            import traceback
            print(f"规划任务时发生异常: {e}")
            print(traceback.format_exc())
            return []
        finally:
            self.satellite_network.is_planed = True  # 设置卫星网络的规划状态为True
            self.time_multiple = self.default_speed_doubling  # 设置时间倍数为默认速度倍数

    # 将数据转换为扁平化的形式
    def flatten_data(self, input_data):
        flattened = []
        for outer_key, satellites in input_data.items():  # outer_key 就是 "task_nb"
            for sat_name, time_windows in satellites.items():  # 遍历 SAT-001, SAT-004...
                for time_range, details in time_windows.items():  # 遍历时间窗口
                    # 构造扁平化字典
                    record = {
                        "task_name": outer_key,  # 新增字段，值为 "task_nb"
                        "satellite": sat_name,
                        "time_range": time_range
                    }
                    # 合并详情字段
                    record.update(details)
                    flattened.append(record)

        return flattened

    # 上注任务
    def distribute_tasks(self, tasks):
        """
        将任务部署到卫星网络中
        """
        # 加锁保护共享 new_tasks 列表（与网络线程 check_execute_tasks 共用同一把 task_lock）
        with self.task_lock:
            for task in tasks:
                if self.satellite_network:
                    # self.satellite_network.deploy_tasks(task)
                    self.satellite_network.new_tasks.append(task)
            # 对新任务列表进行根据优先级排序
            self.satellite_network.new_tasks.sort()

    # 将一个星簇的任务迁移至另外一个星簇，并重新规划任务
    def replan(self, cluster_name1, cluster_name2):
        """
        重新规划任务
        """
        re_tasks = []  # 存储需要重新规划的任务
        print(f"开始重新规划任务，从{cluster_name1}迁移到{cluster_name2}")
        print(f"net_tasks_buffer中的任务数量: {len(self.satellite_network.net_tasks_buffer)}")
        print(f"new_tasks中的任务数量: {len(self.satellite_network.new_tasks)}")

        # 加锁保护共享任务列表/字典的移除与迁移（与网络线程共用同一把 task_lock，RLock 可重入）
        with self.task_lock:
            # 将符合条件的主任务和其子任务从卫星网络中移除
            for main_task in list(self.satellite_network.net_tasks_buffer.values()):
                print(f"检查主任务: {main_task.task_id}, cluster_name={main_task.cluster_name}, status={main_task.status}")
                if main_task.cluster_name == cluster_name1 and main_task.status == "等待执行":
                    print(f"找到需要迁移的主任务: {main_task.task_id}")
                    main_task.cluster_name = cluster_name2

                    del self.satellite_network.net_tasks_buffer[main_task.task_id]
                    self.tasks_buffer[main_task.task_id] = main_task

                    main_task.status = "等待规划"
                    main_task.is_replan = True  # 将主任务标记为重新规划

                    # 更新数据库
                    self.update_task(main_task)
                    # 处理子任务
                    print(f"主任务{main_task.task_id}的子任务数量: {len(main_task.subtasks)}")
                    for subtask in main_task.subtasks:
                        subtask.cluster_name = cluster_name2
                        subtask.earliest_start_time = subtask.earliest_start_time_backup
                        subtask.latest_end_time = subtask.latest_end_time_backup
                        if subtask.assigned_satellite:
                            self.satellite_network.satellites[subtask.assigned_satellite].tasks_len -= 1
                        subtask.assigned_satellite = None
                        subtask.is_replan = True  # 将子任务标记为重新规划
                        self.satellite_network.new_tasks.remove(subtask)
                        re_tasks.append(subtask)  # 将子任务加入重新规划任务列表

            # 处理独立任务
            tasks_to_remove = []
            for task in self.satellite_network.new_tasks:
                # print(
                #     f"检查独立任务: task_id={task.task_id}, cluster_name={task.cluster_name}, parent_task_id={task.parent_task_id}")
                if task.parent_task_id is None and task.cluster_name == cluster_name1 and task.status == "等待执行":
                    print(f"找到需要迁移的独立任务: {task.task_id}")
                    tasks_to_remove.append(task)

            # 先删除任务
            for task in tasks_to_remove:
                # 从等待类表中删除原来的任务
                self.satellite_network.new_tasks.remove(task)
                # 再处理任务
                task.status = "等待规划"
                task.cluster_name = cluster_name2
                task.earliest_start_time = task.earliest_start_time_backup
                task.latest_end_time = task.latest_end_time_backup
                task.is_replan = True  # 将任务标记为重新规划
                if task.assigned_satellite:
                    self.satellite_network.satellites[task.assigned_satellite].tasks_len -= 1
                self.update_task(task)
                print("需要重新规划的独立任务:", task.task_id)
                re_tasks.append(task)

        # 重新规划任务
        if re_tasks:
            for task in re_tasks:
                print(task.task_id)
            print("需要重新规划的任务个数:", len(re_tasks))
            new_tasks = self.planning_tasks(re_tasks, self.model)
            if new_tasks:
                self.distribute_tasks(new_tasks)

    # # 接收结果
    # def receive_results(self):
    #     """
    #     从数据中心接收处理后的数据
    #     """
    #     # 如果有要处理的任务列表
    #     if not self.result_queue.empty():
    #         result_list = self.result_queue.get()
    #         for result in result_list:
    #             # 从未完成任务字典中移除
    #             if result.get("task_id") in self.pending_tasks:
    #                 del self.pending_tasks[result.get("task_id")]
    #             # 添加到已完成任务字典中
    #             self.completed_tasks[result.get("task_id")] = result
    #             # 立即显示结果
    #             print(f"用户收到任务结果: {result}")

    # def export_task_results(self):
    #     """
    #     导出任务结果到文件
    #     未完成任务：覆盖式导出
    #     已完成任务：追加式导出
    #     """
    #     timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    #
    #     # 导出未完成任务（覆盖式）
    #     try:
    #         with open("result/pending_tasks.json", 'w', encoding='utf-8') as f:
    #             json.dump(self.pending_tasks, f, indent=4, ensure_ascii=False)
    #             f.flush()
    #             f.close()
    #         print("未完成任务已导出到 pending_tasks.json")
    #     except Exception as e:
    #         print(f"导出未完成任务时发生错误：{str(e)}")
    #     # 导出已完成任务（追加式）
    #     if self.completed_tasks:
    #         completed_file = "result/completed_tasks.json"
    #         try:
    #             # 读取现有的已完成任务（如果文件存在）
    #             existing_completed = {}
    #             if os.path.exists(completed_file):
    #                 with open(completed_file, 'r', encoding='utf-8') as f:
    #                     try:
    #                         existing_completed = json.load(f)
    #                     except json.JSONDecodeError:
    #                         existing_completed = {}
    #
    #             # 更新已完成任务数据
    #             existing_completed.update(self.completed_tasks)
    #
    #             # 写入更新后的数据
    #             with open(completed_file, 'w', encoding='utf-8') as f:
    #                 json.dump(existing_completed, f, indent=4, ensure_ascii=False)
    #             print("已完成任务已追加到 completed_tasks.json")
    #
    #             # 清空已完成任务字典（因为已经保存到文件中）
    #             self.completed_tasks.clear()
    #
    #         except Exception as e:
    #             print(f"导出已完成任务时发生错误：{str(e)}")

    def update_task(self, task):
        """
        更新任务状态
        :param task:任务对象
        :return:
        """
        with self.app.app_context():
            newtask_model = NewTaskModel.query.filter_by(id=task.task_id).first()
            if not newtask_model:
                print(f"警告：任务 {task.task_id} 不存在，跳过更新")
                return
            newtask_model.status = task.status
            newtask_model.start_time = task.earliest_start_time
            newtask_model.end_time = task.latest_end_time
            newtask_model.cluster_name = task.cluster_name
            newtask_model.assigned_satellite_name = None
            db.session.commit()

    def check_files_of_sat(self):
        """
        检查当前项目的library目录下是否存在TLE.txt和sat_parameters.excel文件
        """
        library_dir = os.path.join(os.getcwd(), 'library')
        # tle_path = os.path.join(library_dir, 'TLE2.txt')
        sat_params_path = os.path.join(library_dir, sat_parms)
        return os.path.exists(sat_params_path)

    def check_files_of_tle(self):
        library_dir = os.path.join(os.getcwd(), 'library')
        tle_path = os.path.join(library_dir, TLE)
        return os.path.exists(tle_path)

    def check_files_of_cluster(self):
        """
        检查当前项目的library目录下是否存在星簇配置文件cluster_parameters.excel
        """
        library_dir = os.path.join(os.getcwd(), 'library')
        cluster_params_path = os.path.join(library_dir, cluster_parms)
        return os.path.exists(cluster_params_path)

    def run(self):
        """
        运控中心线程运行函数
        """
        # 初始化文件策略：library/ 下已存在则直接复用（重启后自动初始化），不再删除；
        # 用户仍可通过「系统设置」上传新文件覆盖（/initTLE、/initFiles 会写同名文件）
        library_dir = os.path.join(os.getcwd(), 'library')
        for fname in (sat_parms, TLE, cluster_parms):
            fpath = os.path.join(library_dir, fname)
            if os.path.exists(fpath):
                print(f"检测到已存在的初始化文件，直接复用: {fpath}")
            else:
                print(f"初始化文件缺失，等待用户上传: {fpath}")

        # 使用保存的app实例
        app = self.app if hasattr(self, 'app') else None

        with app.app_context() if app else nullcontext():
            print("正在等待用户提交TLE初始化文件...")
            while (not self.is_TLE) and (not self.check_files_of_tle()):
                time.sleep(2)
            print("已经提交TLE文件")
            print("正在等待用户提交卫星参数初始化文件...")
            #     如果is_random为False，则代表以文件形式提交数据
            #     判断不存在sat_parameters.excel文件，则等待
            while (not self.is_random) and (not self.check_files_of_sat()):
                time.sleep(2)

            # 创建卫星网络
            try:
                self.satellite_network = SatelliteNetwork(self.network_config, self.sats_config, self, self.is_random,
                                                          app)
            except Exception as e:
                # 初始化失败：打印 traceback 并置标志返回，避免 daemon 线程无声退出
                import traceback
                print(f"卫星网络初始化失败: {e}")
                print(traceback.format_exc())
                self.init_failed = True
                return
            # 启动卫星网络线程
            # 初始化所有的卫星轨迹
            # 初始化所有卫星的轨迹
            print("初始化所有卫星的轨迹")
            for sat_name, sat in self.satellite_network.satellites.items():
                sat.sat_trace_window(self.start_time)
            print("初始化所有卫星的轨迹完成")
            # imageloader = ImageLoader(r"C:\Users\ZhuanZ\Desktop\system_design\RSICD_images")
            network_thread = Thread(target=self.satellite_network.run)
            network_thread.daemon = True
            network_thread.start()

        while True:
            try:
                print("=================待规划任务批数", self.task_queue.qsize(), "================")
                # 所有卫星的状态都初始化一遍才能开始规划任务
                print("是否已经初始化完毕：", self.satellite_network.is_trace)

                if self.satellite_network.is_trace:
                    if self.auto_run:
                        tasks = self.receive_tasks()
                        if tasks:
                            # 规划任务
                            new_tasks = self.planning_tasks(tasks, self.model)
                            # 部署任务
                            if new_tasks:
                                self.distribute_tasks(new_tasks)
                        else:
                            print("暂无任务")
                    else:
                        # 程序控制模式：不自动取队列规划，任务留在队列中等待手动触发（startTask 等）
                        print("程序控制模式：等待手动触发规划")
                    self.now_time = get_now_time_from_start(self.now_time, INTER_VAL_TIME * self.time_multiple)
                print(self.now_time)
                # 同步时间和倍速
                self.satellite_network.now_time = self.now_time
                self.satellite_network.time_multiple = self.time_multiple
            except Exception as e:
                # 后台主循环异常兜底：打印异常并继续，避免线程退出
                import traceback
                print(f"运控主循环发生异常: {e}")
                print(traceback.format_exc())

            time.sleep(INTER_VAL_TIME)  # 控制处理频率
