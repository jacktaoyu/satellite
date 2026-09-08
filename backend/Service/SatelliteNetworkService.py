import math
import os
import random
import pandas as pd
from datetime import timezone, datetime
from contextlib import nullcontext

import socket
import threading
import json
import time
from queue import Queue

from skyfield.api import EarthSatellite, load
import openpyxl

from Service.GroundStation import GroundStation
from database import db
from Service.ClusterService import Cluster
from Service.SatelliteService import Satellite
from model.ClusterModel import ClusterModel, ClusterStarRelation

from model.TaskModel import NewTaskModel, OldTaskModel
from config import NODES, GROUND_STATION, INTER_VAL_TIME


class SatelliteNetwork:
    def __init__(self, Dict_net, sats_config, occ, is_random, app=None):
        self.operation_center = occ
        self.app = app  # 保存app实例
        self.START_TIME = Dict_net['Sat_Walker']['Start_Time']
        self.END_TIME = Dict_net['Sat_Walker']['End_Time']
        self.interval_time = Dict_net['Sat_Walker']['interval_time']
        self.TLE_file = open(Dict_net['Sat_Walker']['TLE_path'], "r")
        self.Num_of_Orbits = Dict_net['Sat_Walker']['Num_of_Orbits']  # 轨道
        self.Num_of_Sats_per_Orbits = Dict_net['Sat_Walker']['Sat_Num_per_Orbit']
        self.clusters_file = Dict_net['Sat_Walker']['Clusters_path']
        self.satellite_params_file = Dict_net['Sat_Walker'].get('Satellite_Params_Path')
        self.satellites = {}  # 存储卫星对象
        self.net_tasks_buffer = {}  # 主任务缓冲区
        self.is_trace = False  # 是否开启已经全部计算完轨迹
        # self.network_state = {}
        self.clusters = []  # 存储星簇对象
        self.orbits = set()  # 存储所有轨道
        self.new_tasks = []  # 新任务队列
        self.running_tasks = []  # 正在执行的任务队列
        self.pause_tasks = []  # 暂停任务队列
        self.results_buffer = Queue()  # 结果缓冲区
        self.is_connect = False  # 是否连接到客户端
        self.ground_stations = []
        self._pending_task_updates = {}  # 批量任务状态更新缓存 {task_id: status}
        for key, value in GROUND_STATION.items():
            self.ground_stations.append(GroundStation(key, value))

        # 时间和倍速信息
        self.now_time = self.START_TIME
        self.time_multiple = 1

        # 是否规划完成，若未规划完成，则停止更新网络
        self.is_planed = True

        try:
            satellite_params = {}
            start_time = self.START_TIME
            end_time = self.END_TIME
            if not is_random:
                # 读取卫星参数Excel文件
                wb = openpyxl.load_workbook(self.satellite_params_file)
                ws = wb.active
                # 从第二行开始读取数据（跳过表头）
                for row in range(2, ws.max_row + 1):
                    sat_name = ws.cell(row=row, column=1).value.strip()
                    if sat_name:
                        satellite_params[sat_name] = {
                            'orbit': int(ws.cell(row=row, column=2).value),
                            'storage': float(ws.cell(row=row, column=3).value),
                            'battery': float(ws.cell(row=row, column=4).value),
                            'sensor_type': ws.cell(row=row, column=5).value,
                            'resolution_capability': float(ws.cell(row=row, column=6).value),
                            'max_roll_angle': float(ws.cell(row=row, column=7).value),
                            'max_pitch_angle': float(ws.cell(row=row, column=8).value),
                            'stable_time': float(ws.cell(row=row, column=9).value),
                            'angle_velocity': float(ws.cell(row=row, column=10).value),
                            'width_of_cloth': float(ws.cell(row=row, column=11).value),
                            'cloud_threshold': float(ws.cell(row=row, column=12).value),
                            'downlink_rate': float(ws.cell(row=row, column=13).value),
                            'sunlight_powers': float(ws.cell(row=row, column=14).value),
                            'maneuver_powers': float(ws.cell(row=row, column=15).value),
                            'imaging_powers': float(ws.cell(row=row, column=16).value),
                            'eclipse_powers': float(ws.cell(row=row, column=17).value),
                            'start_time': start_time,
                            'end_time': end_time
                        }
                wb.close()
            # 读取TLE文件创建卫星
            data_list = self.TLE_file.readlines()

            orbit_info_en = {}  # 新增：用于统计每个轨道的载荷-分辨率信息
            # 创建卫星对象
            for num in range(0, len(data_list), 3):
                if num + 2 < len(data_list) and data_list[num].strip() != '':
                    sat_name = data_list[num].strip()
                    ts = load.timescale()
                    tle_line1 = data_list[num + 1].strip()
                    tle_line2 = data_list[num + 2].strip()

                    if is_random:  # 如果不使用excle文件获得参数，则使用前端页面指定的参数
                        print(sats_config)
                        if sat_name in sats_config:
                            if sat_name == 'Sat_10_0':
                                print("在随机模式下，Sat_10_0被选中")
                            satellite_params[sat_name] = {
                                'orbit': int(sat_name.split('_')[1]),
                                'storage': sats_config[sat_name].get('max_storage'),
                                'battery': sats_config[sat_name].get('max_battery'),
                                'sensor_type': sats_config[sat_name].get('sensor_type'),
                                'resolution_capability': sats_config[sat_name].get('resolution'),
                                'max_roll_angle': sats_config[sat_name].get('side_swing_angle'),
                                'max_pitch_angle': sats_config[sat_name].get('pitch_angle'),
                                'stable_time': sats_config[sat_name].get('stable_time'),
                                'angle_velocity': sats_config[sat_name].get('angle_velocity'),
                                'width_of_cloth': sats_config[sat_name].get('width_of_cloth'),
                                'cloud_threshold': sats_config[sat_name].get('thickness_threshold'),
                                'start_time': start_time,
                                'end_time': end_time,
                                'imaging_powers': sats_config[sat_name].get('imaging_powers'),
                                'sunlight_powers': sats_config[sat_name].get('sunlight_powers'),
                                'eclipse_powers': sats_config[sat_name].get('eclipse_powers'),
                                'maneuver_powers': sats_config[sat_name].get('maneuver_powers'),
                                'downlink_rate': sats_config[sat_name].get('downlink_rate')
                            }
                        else:
                            sensor_type = random.choice(['optical', 'SAR', 'infrared'])
                            satellite_params[sat_name] = {
                                'orbit': int(sat_name.split('_')[1]),
                                'storage': Dict_net[sensor_type].get('max_storage'),
                                'battery': Dict_net[sensor_type].get('max_battery'),
                                'sensor_type': sensor_type,
                                # 'resolution_capability': random.choice([1.0, 1.5, 2.0]),
                                'resolution_capability': Dict_net[sensor_type].get('resolution'),
                                'max_roll_angle': Dict_net[sensor_type].get('side_swing_angle'),
                                'max_pitch_angle': Dict_net[sensor_type].get('pitch_angle'),
                                'stable_time': Dict_net[sensor_type].get('stable_time'),
                                'angle_velocity': Dict_net[sensor_type].get('angle_velocity'),
                                'width_of_cloth': Dict_net[sensor_type].get('width_of_cloth'),
                                'cloud_threshold': Dict_net[sensor_type].get('thickness_threshold'),
                                'start_time': start_time,
                                'end_time': end_time,
                                'imaging_powers': Dict_net[sensor_type].get('imaging_powers'),
                                'sunlight_powers': Dict_net[sensor_type].get('sunlight_powers'),
                                'eclipse_powers': Dict_net[sensor_type].get('eclipse_powers'),
                                'maneuver_powers': Dict_net[sensor_type].get('maneuver_powers'),
                                'downlink_rate': Dict_net[sensor_type].get('downlink_rate')
                            }
                        # satellite_params[sat_name] = {
                        #     'orbit': int(sat_name.split('_')[1]),
                        #     'storage': Dict_net[sensor_type].get('max_storage'),
                        #     'battery': Dict_net[sensor_type].get('max_battery'),
                        #     'sensor_type': sensor_type,
                        #     'resolution_capability': random.choice([1.0, 1.5, 2.0]),
                        #     # 'resolution_capability': Dict_net[sensor_type].get('resolution'),
                        #     'max_roll_angle': Dict_net[sensor_type].get('side_swing_angle'),
                        #     'max_pitch_angle': Dict_net[sensor_type].get('pitch_angle'),
                        #     'stable_time': Dict_net[sensor_type].get('stable_time'),
                        #     'angle_velocity': Dict_net[sensor_type].get('angle_velocity'),
                        #     'width_of_cloth': random.randint(12, Dict_net[sensor_type].get('width_of_cloth')),
                        #     'cloud_threshold': Dict_net[sensor_type].get('thickness_threshold'),
                        #     'start_time': start_time,
                        #     'end_time': end_time
                        # }
                    # 获取该卫星的参数
                    params = satellite_params.get(sat_name)
                    orbit = params.get('orbit')
                    sensor_type = params.get('sensor_type')  # 传感器类型
                    # resolution = params.get('resolution_capability', 0.5),  # 分辨率能力
                    resolution = params.get('resolution_capability', 0.5)  # 分辨率能力
                    if sat_name == 'Sat_10_0':
                        print("最终的参数")
                        print(params.get('storage', 500))
                    # 创建卫星对象
                    sat = Satellite(
                        sat_id=num // 3 + 1,  # 卫星ID
                        sat_name=sat_name,  # 卫星名称
                        tle_line1=tle_line1,  # TLE行1
                        tle_line2=tle_line2,  # TLE行2
                        sat_model=EarthSatellite(tle_line1, tle_line2, sat_name, ts),  # 卫星模型
                        orbit=params.get('orbit'),  # 轨道
                        storage=params.get('storage', 500),  # 存储容量
                        battery=params.get('battery', 5000),  # 电池容量
                        sensor_type=params.get('sensor_type'),  # 传感器类型
                        resolution_capability=params.get('resolution_capability', 0.5),  # 分辨率能力
                        max_roll_angle=params.get('max_roll_angle', 45),  # 最大侧摆角
                        max_pitch_angle=params.get('max_pitch_angle', 45),  # 最大俯仰角
                        stable_time=params.get('stable_time', 10),  # 稳定时间
                        angle_velocity=params.get('angle_velocity', 5.0),  # 角度速度
                        cloud_threshold=params.get('cloud_threshold', 800),  # 云检测阈值
                        width_of_cloth=params.get('width_of_cloth', 50),  # 幅宽
                        simulation_start_time=params.get('start_time').replace(tzinfo=timezone.utc),  # 仿真开始时间
                        end_time=params.get('end_time').replace(tzinfo=timezone.utc),  # 仿真结束时间
                        client=(num // 3 + 1) % NODES,  # 使用模运算确保client值在0到NODES-1之间
                        imaging_powers=params.get('imaging_powers', 50),  # 成像功率
                        sunlight_powers=params.get('sunlight_powers', 30),  # 太阳能功率
                        eclipse_powers=params.get('eclipse_powers', 8),  # 遮挡功率
                        maneuver_powers=params.get('maneuver_powers', 500),  # 机动功率
                        downlink_rate=params.get('downlink_rate', 4),  # 下行速率
                    )
                    self.satellites[sat_name] = sat  # 将卫星添加到字典中
                    self.orbits.add(sat.orbit)  # 将轨道添加到集合中

                    # 英文转中文
                    sensor_type_map = {'optical': '光学', 'infrared': '红外', 'SAR': 'SAR'}
                    sensor_type_cn = sensor_type_map.get(sensor_type, sensor_type)

                    # 统计到字典
                    if orbit not in orbit_info_en:
                        orbit_info_en[orbit] = {}
                    if sensor_type_cn not in orbit_info_en[orbit]:
                        orbit_info_en[orbit][sensor_type_cn] = set()
                    orbit_info_en[orbit][sensor_type_cn].add(resolution)
                    # # 将卫星数据写入Excel文件
                    # satellite_data = {
                    #     '卫星名': sat.sat_name,
                    #     '轨道': sat.orbit,
                    #     '存储容量(GB)': sat.storage,
                    #     '电池容量(mAh)': sat.battery,
                    #     '传感器类型': sat.star_payload,
                    #     '分辨率(m)': sat.resolution_capability,
                    #     '最大侧摆角(度)': sat.side_swing_angle_Max,
                    #     '最大俯仰角(度)': sat.pitch_angle_Max,
                    #     '稳定时间(s)': sat.stable_time,
                    #     '角度速度(度/s)': sat.angle_velocity,
                    #     '幅宽(km)': sat.width_of_cloth,
                    #     '云检测阈值': sat.cloud_threshold
                    # }
                    #
                    # # 创建DataFrame并写入Excel
                    # df = pd.DataFrame([satellite_data])
                    # excel_path = os.path.join('output_plans', 'satellite_info.xlsx')
                    #
                    # # 如果文件已存在，追加数据
                    # if os.path.exists(excel_path):
                    #     with pd.ExcelWriter(excel_path, mode='a', engine='openpyxl', if_sheet_exists='overlay') as writer:
                    #         df.to_excel(writer, index=False, header=False, startrow=writer.sheets['Sheet1'].max_row)
                    # else:
                    #     # 如果文件不存在，创建新文件
                    #     df.to_excel(excel_path, index=False)

                # 统计结束后，格式化输出
            orbit_info = {}
            for orbit, payloads in orbit_info_en.items():
                parts = []
                for payload, resolutions in payloads.items():
                    res_str = ",".join([f"{r}" for r in sorted(resolutions)])
                    parts.append(f"{payload}:{res_str}")
                orbit_info[orbit] = f"第{orbit}轨道" + "|".join(parts)

            self.orbit_info = orbit_info
            print("orbit_info：", orbit_info)

        except Exception as e:
            print(f"读取卫星参数或创建卫星对象时发生错误：{str(e)}")
            import traceback
            print(traceback.format_exc())
        finally:
            self.TLE_file.close()

        # 创建星簇 
        self.create_clusters(self.clusters_file)

        # 尝试不同的端口
        port = 9999
        max_retries = 10
        for _ in range(max_retries):
            try:
                self.server_socket = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
                # 保持 0.0.0.0：SatClient.exe 实物客户端需要从局域网连接本 socket 服务，不能收窄为 127.0.0.1
                # 鉴权：客户端连接后需先发送 {"token": ...} 握手消息（见 _verify_client_token）
                self.server_socket.bind(('0.0.0.0', port))
                self.server_socket.listen(NODES)
                print(f"成功绑定端口 {port}")
                break
            except OSError:
                print(f"端口 {port} 被占用，尝试下一个端口")
                self.server_socket.close()  # 关闭旧socket，避免文件描述符泄漏
                port += 1
        else:
            raise RuntimeError(f"无法找到可用端口，请检查 {port - max_retries} 到 {port - 1} 端口的占用情况")
        # # 新增 socket 服务器
        # self.server_socket = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        # self.server_socket.bind(('0.0.0.0', 9999))
        # self.server_socket.listen(NODES)

        # socket 握手鉴权 token：环境变量 SOCKET_AUTH_TOKEN 可覆盖，未设置时使用开发默认值
        self.auth_token = os.environ.get('SOCKET_AUTH_TOKEN')
        if not self.auth_token:
            self.auth_token = 'dev-satellite-token'
            print("警告: 未设置环境变量 SOCKET_AUTH_TOKEN，socket 服务使用开发默认 token，生产环境请务必配置")

        # 客户端连接列表
        self.client_sockets = []
        # 客户端发送线程列表
        self.client_send_threads = []
        # 客户端接收线程列表
        self.client_recv_threads = []

    def _start_socket_server(self):
        """启动 socket 服务器：非阻塞 accept，主循环内逐步接收客户端连接。

        原先此处 while 循环永久阻塞等待 NODES 个客户端，无客户端时卫星网络线程
        无法进入主循环，is_trace 永远为 False，任务规划不会启动；改为每次主循环
        尝试一次带超时的 accept，无客户端时规划/仿真照常运行，客户端可随时接入。
        """
        self.server_socket.settimeout(1.0)  # accept 超时后返回主循环，避免阻塞规划
        print("socket 服务已就绪（非阻塞模式），客户端可随时连接")

    def _try_accept_client(self):
        """尝试接收一个客户端连接（非阻塞）；无待处理连接时直接返回"""
        if len(self.client_sockets) >= NODES:
            return
        try:
            client_socket, client_address = self.server_socket.accept()
        except (socket.timeout, BlockingIOError):
            return
        except OSError:
            return
        if not self._verify_client_token(client_socket, client_address):
            client_socket.close()
            return
        print(f"客户端 {client_address[0]}:{client_address[1]} 已连接")
        self.client_sockets.append(client_socket)
        # 为每个客户端创建接收线程
        recv_thread = threading.Thread(target=self._receive_results, args=(client_socket,))
        self.client_recv_threads.append(recv_thread)
        recv_thread.start()
        if len(self.client_sockets) >= NODES:
            print("所有客户端已连接")
            self.is_connect = True

    def _verify_client_token(self, client_socket, client_address):
        """校验客户端第一条握手消息中的 token，失败返回 False"""
        try:
            client_socket.settimeout(10)  # 握手超时保护，避免永久阻塞 accept 循环
            data = client_socket.recv(1024)
            client_socket.settimeout(None)
            msg = json.loads(data.decode().strip())
            if msg.get('token') == self.auth_token:
                return True
            print(f"客户端 {client_address[0]}:{client_address[1]} 鉴权失败，关闭连接")
        except Exception as e:
            print(f"客户端 {client_address[0]}:{client_address[1]} 鉴权握手异常: {e}，关闭连接")
        return False

    def _send_tasks(self, satellite, task, execute_time=None):
        """发送任务到客户端"""
        if satellite.client >= len(self.client_sockets):
            print(
                f"Error: Invalid client index {satellite.client}, client_sockets length is {len(self.client_sockets)}")
            return

        msg = {
            "task_id": task.task_id,
            "execute_time": execute_time if execute_time is not None else task.execution_time
        }
        # 确保发送完整JSON数据，添加换行符作为消息分隔符
        data = json.dumps(msg) + "\n"
        try:
            print("发送任务给客户端：", task.task_id, "卫星的客户端是:", satellite.client,
                  self.client_sockets[satellite.client])
            self.client_sockets[satellite.client].sendall(data.encode())
        except Exception as e:
            print(f"发送失败: {e}")

    def _receive_results(self, client_socket):
        """接收客户端返回的结果"""
        buffer = b""
        while True:
            try:
                data = client_socket.recv(1024)
                if not data:
                    break

                buffer += data
                # 按换行符分割消息
                while b'\n' in buffer:
                    msg_str, buffer = buffer.split(b'\n', 1)
                    try:
                        # 解析JSON消息
                        msg = json.loads(msg_str.decode())
                        # 将结果放到结果缓冲区
                        self.results_buffer.put(msg)
                    except json.JSONDecodeError as e:
                        print(f"JSON解析错误: {e}")
                        continue
            except Exception as e:
                print(f"接收结果出错: {str(e)}")
                break
        # 客户端断连清理
        self._handle_client_disconnect(client_socket)

    def _handle_client_disconnect(self, client_socket):
        """客户端断连清理：关闭 socket、移出连接列表，并将其上“正在执行”的任务重置回等待执行"""
        try:
            client_socket.close()
        except Exception:
            pass
        if client_socket not in self.client_sockets:
            return
        client_index = self.client_sockets.index(client_socket)
        self.client_sockets.remove(client_socket)
        print(f"客户端连接已断开并清理，剩余客户端数: {len(self.client_sockets)}")
        # 重置该客户端上处于“正在执行”状态且未收到结果的任务，避免任务永久卡死
        for task in self.running_tasks[:]:
            satellite = self.satellites.get(task.assigned_satellite)
            if satellite is None or satellite.client != client_index or task.status != "正在执行":
                continue
            self.running_tasks.remove(task)
            task.status = "等待执行"
            if satellite.running_task == task.task_id:
                satellite.running_task = None
                satellite.status = "FREE"
            # 重新加入待执行队列（与 check_execute_tasks 共用同一把锁）
            lock = self.operation_center.task_lock if self.operation_center else nullcontext()
            with lock:
                self.new_tasks.append(task)
            self.update_task(task)
            print(f"任务 {task.task_id} 因客户端断连重置为等待执行")

    def _create_cluster_in_db(self, cluster, orbit_str, payload_resolution_map):
        """在数据库中创建星簇记录"""
        with self.app.app_context():
            payload_resolution_map_str = str(payload_resolution_map).strip('{}')
            cluster_model = ClusterModel(
                name=cluster.name,
                orbits=orbit_str,
                satellite_count=cluster.satellite_count,
                payload_resolution=payload_resolution_map_str,
                status=True
            )
            db.session.add(cluster_model)
            db.session.commit()
            return cluster_model.id

    # 根据星簇文件创建星簇
    def create_clusters(self, excel_path):
        """
        创建卫星星簇
        Args:
            excel_path: 包含星簇参数的Excel文件路径
        """
        if os.path.exists(excel_path):
            try:
                # 加载Excel文件
                wb = openpyxl.load_workbook(excel_path)
                ws = wb.active  # 使用第一个sheet
                # 清空现有的星簇
                self.clusters = []

                # 从第二行开始读取星簇要求（第一行是表头）
                for row in range(2, ws.max_row + 1):
                    cluster_name = ws.cell(row=row, column=1).value
                    if not cluster_name:  # 跳过空行
                        continue

                    # 创建新的星簇对象
                    cluster = Cluster(cluster_name)
                    payload_resolution_map = {}

                    # 读取要求
                    payload_resolution_value = ws.cell(row=row, column=3).value  # 第三列为Payload-Resolution Map
                    # 新格式：载荷类型及其对应的分辨率要求
                    if payload_resolution_value:
                        sensor_type_parts = str(payload_resolution_value).split('|')
                        for part in sensor_type_parts:
                            if ':' in part:
                                sensor_type, resolutions_str = part.strip().split(':')
                                resolutions = [float(r.strip()) for r in resolutions_str.split(',')]
                                payload_resolution_map[sensor_type] = resolutions
                            else:
                                # 向后兼容：如果没有指定分辨率，则认为可以接受任何分辨率
                                sensor_type = part.strip()
                                payload_resolution_map[sensor_type] = []

                    # 解析轨道要求
                    orbit_str = ws.cell(row=row, column=2).value
                    orbit_list = [int(o.strip()) for o in str(orbit_str).split(',')] if orbit_str else []
                    cluster.orbits = set(orbit_list)  # 将列表转换为集合
                    # 遍历所有卫星，找出满足条件的卫星
                    for sat_name, satellite in self.satellites.items():
                        # 检查轨道要求
                        if orbit_list:
                            orbit_num = int(sat_name.split('_')[1])
                            if orbit_num not in orbit_list:
                                continue

                        # 检查载荷类型和分辨率要求
                        sensor_type = satellite.star_payload
                        resolution = satellite.resolution_capability

                        # 检查载荷类型是否在要求列表中
                        if sensor_type not in payload_resolution_map:
                            continue

                        # 如果为此载荷类型指定了分辨率要求，则检查分辨率
                        if payload_resolution_map.get(sensor_type) and len(
                                payload_resolution_map[sensor_type]) > 0 and resolution not in payload_resolution_map[
                            sensor_type]:
                            continue

                        # 如果卫星满足所有要求，将其添加到星簇中
                        cluster.append_sta(satellite)
                        # 星簇名加入到卫星的星簇表中
                        satellite.belong_cluster.append(cluster.name)

                    # 将星簇添加到列表中
                    self.clusters.append(cluster)
                    print(cluster.name, cluster.payload_resolution_map)
                    # print('星簇分辨率', cluster.resolution)
                    # 将星簇添加到数据库中
                    cluster.cluster_id = self._create_cluster_in_db(cluster, orbit_str, payload_resolution_value)
                    self.add_sat_cluster_relation(cluster)
                    # 更新Excel文件中的卫星数量和卫星名称
                    ws.cell(row=row, column=4, value=cluster.satellite_count)  # Satellite Count
                    ws.cell(row=row, column=5, value=str([sat.sat_name for sat in cluster.stars]))  # Satellite Names

                # 保存更新后的Excel文件
                wb.save(excel_path)
                print(f"已成功创建{len(self.clusters)}个星簇")

            except FileNotFoundError:
                print(f"未找到星簇参数文件：{excel_path}")
            except Exception as e:
                print(f"创建星簇时发生错误：{str(e)}")
                import traceback
                print(traceback.format_exc())  # 打印详细的错误信息
        else:
            print("未提供星簇参数文件路径,采用默认星簇分类")
            # 清空现有的星簇
            self.clusters = []

            # 创建一个字典来存储轨道-载荷-分辨率组合对应的星簇
            cluster_map = {}

            # 遍历所有卫星
            for sat_name, satellite in self.satellites.items():
                # 从卫星名称中提取轨道编号
                orbit_num = int(sat_name.split('_')[1])

                # 创建星簇的唯一标识符
                cluster_key = (orbit_num, satellite.star_payload, satellite.resolution_capability)

                # 如果这个组合还没有对应的星簇，创建一个新的星簇
                if cluster_key not in cluster_map:
                    cluster_name = f"Cluster_{orbit_num}_{satellite.star_payload}_{satellite.resolution_capability}"
                    cluster = Cluster(cluster_name)
                    # cluster_orbits = [orbit_num]
                    cluster_map[cluster_key] = cluster
                    self.clusters.append(cluster)

                # 将卫星添加到对应的星簇中
                cluster_map[cluster_key].append_sta(satellite)
                # 星簇名加入到卫星的星簇表中
                satellite.belong_cluster.append(cluster_map[cluster_key].name)
            # 将星簇添加到数据库中
            for cluster in self.clusters:
                orbit_str = str(cluster.orbits).strip('{}')
                payload_resolution_map = self.turn_map_to_str(cluster.payload_resolution_map)
                cluster.cluster_id = self._create_cluster_in_db(cluster, orbit_str, payload_resolution_map)
                self.add_sat_cluster_relation(cluster)

        # 添加一个默认星簇，包含所有卫星
        cluster = Cluster("All_Sat")
        for sat_name, satellite in self.satellites.items():
            cluster.append_sta(satellite)
            satellite.belong_cluster.append(cluster.name)
        orbit_str = str(cluster.orbits).strip('{}')
        payload_resolution_map = self.turn_map_to_str(cluster.payload_resolution_map)
        cluster.cluster_id = self._create_cluster_in_db(cluster, orbit_str, payload_resolution_map)
        self.add_sat_cluster_relation(cluster)
        self.clusters.append(cluster)
        print(f"已成功创建{len(self.clusters)}个默认星簇")

    def add_sat_cluster_relation(self, cluster):
        """
        添加星簇和卫星的关系
        :param cluster:
        :return:
        """
        with self.app.app_context():
            for star in cluster.stars:
                cluster_star_relation_model = ClusterStarRelation(sat_name=star.sat_name,
                                                                  cluster_id=cluster.cluster_id)
                db.session.add(cluster_star_relation_model)
            db.session.commit()  # 批量添加后一次提交

    def turn_map_to_str(self, my_map):
        """
        将星簇的载荷：分辨率字典转换成字符串形式
        """
        strs = ""
        for key, value in my_map.items():
            v = ""
            for i in value:
                v += str(i) if v == "" else f",{i}"
            strs += f"{key}:{v}" if strs == "" else f"|{key}:{v}"
        return strs

    # 根据星簇id删除星簇
    def delete_cluster_by_cluster_id(self, cluster_id):
        """
        根据星簇id删除星簇
        """
        with self.app.app_context():
            clusterModel = ClusterModel.query.filter_by(id=cluster_id).first()
            if clusterModel:
                db.session.delete(clusterModel)
                relations = ClusterStarRelation.query.filter_by(cluster_id=cluster_id).all()
                for relation in relations:
                    db.session.delete(relation)
                    db.session.commit()
                for cluster in self.clusters:
                    if cluster.cluster_id == cluster_id:
                        self.clusters.remove(cluster)
                        for star in cluster.stars:
                            star.belong_cluster.remove(cluster.name)
                        break
                return True
            else:
                return False

    # 构建单个星簇
    def creat_single_cluster(self, name, orbit, payload_resolution_map):
        """
        构建单个星簇
        """
        # 添加一个星簇
        cluster = Cluster(name)
        orbits_list = [int(o.strip()) for o in str(orbit).strip('[]').split(',')] if orbit else []
        cluster.orbits = set(orbits_list)  # 将列表转换为集合
        map = {}
        if payload_resolution_map:
            sensor_type_parts = str(payload_resolution_map).strip('').split('|')
            for part in sensor_type_parts:
                if ':' in part:
                    sensor_type, resolutions_str = part.strip().split(':')
                    resolutions = [float(r.strip()) for r in resolutions_str.split(',')]
                    map[sensor_type] = resolutions
                else:
                    # 向后兼容：如果没有指定分辨率，则认为可以接受任何分辨率
                    sensor_type = part.strip()
                    map[sensor_type] = []
            # 遍历所有卫星，找出满足条件的卫星
            for sat_name, satellite in self.satellites.items():
                # 检查轨道要求
                if orbits_list:
                    orbit_num = int(sat_name.split('_')[1])
                    if orbit_num not in orbits_list:
                        continue

                # 检查载荷类型和分辨率要求
                sensor_type = satellite.star_payload
                resolution = satellite.resolution_capability

                # 检查载荷类型是否在要求列表中
                if sensor_type not in map:
                    continue

                # 如果为此载荷类型指定了分辨率要求，则检查分辨率
                if map.get(sensor_type) and len(
                        map[sensor_type]) > 0 and resolution not in map[sensor_type]:
                    continue

                # 如果卫星满足所有要求，将其添加到星簇中
                cluster.append_sta(satellite)
                satellite.belong_cluster.append(cluster.name)

        # 将星簇添加到数据库中
        cluster.cluster_id = self._create_cluster_in_db(cluster, orbit, payload_resolution_map)
        self.add_sat_cluster_relation(cluster)
        # 将星簇添加到列表中
        self.clusters.append(cluster)
        return True

    # 更改星簇
    def update_cluster(self, cluster_id, name, orbit, payload_resolution_map):
        """
        更改星簇
        """
        # 删除旧的星簇
        for cluster in self.clusters:
            if cluster.cluster_id == cluster_id:
                # 删除星簇和卫星的所属星簇名
                self.clusters.remove(cluster)
                for satellite in cluster.stars:
                    satellite.belong_cluster.remove(cluster.name)
                # 删除星簇和卫星的关系
                with self.app.app_context():
                    relations = ClusterStarRelation.query.filter_by(cluster_id=cluster_id).all()
                    for relation in relations:
                        db.session.delete(relation)
                    db.session.commit()  # 批量删除后一次提交
                break
        # 创建新的星簇
        cluster = Cluster(name)
        map = {}
        orbits_list = [int(o.strip()) for o in str(orbit).strip('[]').split(',')] if orbit else []
        cluster.orbits = set(orbits_list)  # 将列表转换为集合
        if payload_resolution_map:
            sensor_type_parts = str(payload_resolution_map).strip('').split('|')
            for part in sensor_type_parts:
                if ':' in part:
                    sensor_type, resolutions_str = part.strip().split(':')
                    resolutions = [float(r.strip()) for r in resolutions_str.split(',')]
                    map[sensor_type] = resolutions
                else:
                    # 向后兼容：如果没有指定分辨率，则认为可以接受任何分辨率
                    sensor_type = part.strip()
                    map[sensor_type] = []
            # 遍历所有卫星，找出满足条件的卫星
            for sat_name, satellite in self.satellites.items():
                # 检查轨道要求
                if orbits_list:
                    if satellite.orbit not in orbits_list:
                        continue

                # 检查载荷类型和分辨率要求
                sensor_type = satellite.star_payload
                resolution = satellite.resolution_capability

                # 检查载荷类型是否在要求列表中
                if sensor_type not in map:
                    continue

                # 如果为此载荷类型指定了分辨率要求，则检查分辨率
                if map.get(sensor_type) and len(map[sensor_type]) > 0 and resolution not in map[sensor_type]:
                    continue

                # 如果卫星满足所有要求，将其添加到星簇中
                cluster.append_sta(satellite)
                satellite.belong_cluster.append(cluster.name)
        cluster.cluster_id = cluster_id
        self.add_sat_cluster_relation(cluster)
        # 将星簇添加到列表中
        self.clusters.append(cluster)
        with self.app.app_context():
            # 更新数据库
            cluster_model = ClusterModel.query.filter_by(id=cluster_id).first()
            if cluster_model is None:
                print(f"更新星簇失败：数据库中不存在 cluster_id={cluster_id} 的星簇")
                return False
            cluster_model.name = name
            # 将orbits列表转换为字符串
            cluster_model.orbits = str(orbits_list).strip('[]')
            cluster_model.satellite_count = len(cluster.stars)
            cluster_model.payload_resolution = payload_resolution_map
            db.session.commit()
        return True

    def delete_and_recreate_all_clusters(self, excel_path):
        """
        删除并根据文件重新创建所有星簇
        """
        with self.app.app_context():
            # 先清空ClusterStarRelation表
            relations = ClusterStarRelation.query.all()
            for relation in relations:
                db.session.delete(relation)

            # 再清空ClusterModel表
            clusters = ClusterModel.query.all()
            for cluster in clusters:
                db.session.delete(cluster)

            db.session.commit()

            for sat in self.satellites.values():
                sat.belong_cluster = []

        # 重新创建所有星簇
        self.create_clusters(excel_path)

    # 更新网络状态
    def update_net_state(self, now_time, time_multiple):
        # 更新网络状态
        print("更新网络状态")
        for sat_name, satellite in self.satellites.items():
            satellite.update_state(now_time, time_multiple, self.ground_stations)
        print("更新网络状态完成")
        # for ground_station in self.ground_stations:
        #     print(len(ground_station.connecting_satellite), ground_station.connecting_satellite)
        # 表明所有的卫星已经计算了一遍轨迹
        self.is_trace = True

    def collect_results(self):
        """
        收集结果
        :return:
        """
        tasks_to_migrate = []  # 收集需要迁移到旧表的任务
        tasks_to_remove = []   # 收集需要从 running_tasks 移除的任务

        while not self.results_buffer.empty():
            result = self.results_buffer.get()
            print(result)
            # 根据 task_id 查找 running_tasks 中的任务（遍历副本避免修改中遍历）
            for task in self.running_tasks[:]:
                if task.task_id == result["task_id"]:
                    tasks_to_remove.append(task)
                    if task.task_id == self.satellites[task.assigned_satellite].running_task:
                        self.satellites[task.assigned_satellite].running_task = None
                        self.satellites[task.assigned_satellite].status = "FREE"

                    x, y, z = self.satellites[task.assigned_satellite].position
                    # 计算高度
                    height = math.sqrt(x ** 2 + y ** 2 + z ** 2) - 6371
                    # 如果是子任务，需要判断其父任务是否还有其他子任务未完成
                    locations = None
                    if task.parent_task_id:
                        main_task = self.net_tasks_buffer[task.parent_task_id]
                        if task.task_id in main_task.subtask_ids:
                            main_task.subtask_ids.remove(task.task_id)
                        # 如果主任务没有其他子任务，则将其移除
                        if len(main_task.subtask_ids) == 0:
                            task = self.net_tasks_buffer.pop(main_task.task_id)
                            locations = [subtask.target_location for subtask in main_task.subtasks]
                        else:
                            break
                    task.status = "Success"
                    tasks_to_migrate.append({
                        'task': task,
                        'height': height,
                        'locations': locations
                    })
                    break

        # 批量从 running_tasks 移除
        for task in tasks_to_remove:
            if task in self.running_tasks:
                self.running_tasks.remove(task)

        # 批量数据库操作：一次 app_context，一次 commit
        if tasks_to_migrate:
            with self.app.app_context():
                for item in tasks_to_migrate:
                    task = item['task']
                    height = item['height']
                    locations = item['locations']
                    newtask_model = NewTaskModel.query.filter_by(id=task.task_id).first()
                    if newtask_model:
                        oldtask_model = OldTaskModel(
                            id=newtask_model.id,
                            task_name=newtask_model.task_name,
                            priority=newtask_model.priority,
                            is_emergency=newtask_model.is_emergency,
                            sensor_type=newtask_model.sensor_type,
                            task_type=newtask_model.task_type,
                            resolution=newtask_model.resolution,
                            start_time=newtask_model.start_time,
                            end_time=newtask_model.end_time,
                            appoint_time=newtask_model.appoint_time,
                            cloud_thickness=newtask_model.cloud_thickness,
                            target_location=newtask_model.target_location,
                            assigned_satellite_name=newtask_model.assigned_satellite_name,
                            status=task.status,
                            friend_task_id=newtask_model.friend_task_id,
                            is_photo=False,
                            path=None,
                            height=height,
                            sub_length=len(locations) if task.subtasks else None,
                            sub_locations=str(locations) if task.subtasks else None,
                            comment=newtask_model.comment
                        )
                        db.session.delete(newtask_model)
                        db.session.add(oldtask_model)
                db.session.commit()

    def check_execute_tasks(self, now_time, time_multiple):
        """
        检查所有卫星的任务队列并执行任务
        """
        # 与 Flask 请求线程（pause_task/start_task/delete_task）和 OCC 线程（distribute_tasks/replan）
        # 共用同一把锁（OCC 的 task_lock），保护共享的 new_tasks 列表
        lock = self.operation_center.task_lock if self.operation_center else nullcontext()
        with lock:
            if not (self.new_tasks and len(self.client_sockets) > 0):
                return
            # 锁内拍快照，遍历快照执行（循环体内有 socket 发送等耗时操作，避免长时间持锁）
            tasks_snapshot = list(self.new_tasks)
        executed_tasks = []
        for task in tasks_snapshot:
            # 判断是否到达任务的执行时间
            if (task.earliest_start_time.replace(
                    tzinfo=None) - now_time).total_seconds() > INTER_VAL_TIME * time_multiple:
                continue
            satellite = self.satellites[task.assigned_satellite]
            if satellite.is_available:
                # 执行任务
                self.running_tasks.append(task)
                executed_tasks.append(task)
                # 执行时间除以加速倍数（局部变量计算，不修改 task 对象本身，避免断连重发后重复除）
                execute_time = task.execution_time / time_multiple
                # 发送到卫星实物节点执行
                self._send_tasks(satellite, task, execute_time)
                satellite.status = "RUNNING"
                satellite.battery = task.battery_after_task
                # 更新卫星状态
                satellite.storage -= task.current_storage  # 更新存储
                satellite.side_swing_angle = task.target_attitude[1]  # 更新角度
                satellite.pitch_angle = task.target_attitude[2]  # 更新姿态
                satellite.tasks_len -= 1  # 更新任务数量

                task.status = "正在执行"
                if task.parent_task_id and task.parent_task_id in self.net_tasks_buffer:
                    self.net_tasks_buffer[task.parent_task_id].status = "正在执行"
                    satellite.running_task = task.task_id  # 更新任务ID
                    self.update_task(self.net_tasks_buffer[task.parent_task_id], batch=True)  # 批量更新
                else:
                    satellite.running_task = task.task_id  # 更新任务ID
                    self.update_task(task, batch=True)
        # 批量移除已执行任务，避免循环内多次 O(n) remove
        if executed_tasks:
            executed_set = set(executed_tasks)
            with lock:  # 重建 new_tasks 需与写线程互斥
                self.new_tasks = [t for t in self.new_tasks if t not in executed_set]

    # 从新任务模型中删除任务，并将其放到旧表中
    def delete_and_add_task(self, task_id):
        """
        从新任务模型中删除任务，并将其放到旧表中
        :param task_id: 任务ID
        :return:
        """
        with self.app.app_context():
            newtask_model = NewTaskModel.query.filter_by(id=task_id).first()
            if newtask_model:
                oldtask_model = OldTaskModel(
                    id=newtask_model.id,
                    task_name=newtask_model.task_name,
                    priority=newtask_model.priority,
                    is_emergency=newtask_model.is_emergency,
                    sensor_type=newtask_model.sensor_type,
                    task_type=newtask_model.task_type,
                    resolution=newtask_model.resolution,
                    start_time=newtask_model.start_time,
                    end_time=newtask_model.end_time,
                    appoint_time=newtask_model.appoint_time,
                    target_location=newtask_model.target_location,
                    assigned_satellite_name=newtask_model.assigned_satellite_name,
                    status="Failed",
                    is_photo=False,
                    path=None,
                    cloud_thickness=newtask_model.cloud_thickness
                )
                db.session.delete(newtask_model)
                db.session.add(oldtask_model)
                db.session.commit()

    def update_task(self, task, batch=False):
        """
        更新任务状态
        :param task: 任务对象
        :param batch: 若为True，仅将变更加入待处理队列，不立即commit
        :return:
        """
        if batch:
            self._pending_task_updates[task.task_id] = task.status
            return
        with self.app.app_context():
            newtask_model = NewTaskModel.query.filter_by(id=task.task_id).first()
            if newtask_model:
                newtask_model.status = task.status
                db.session.commit()

    def flush_task_updates(self):
        """
        批量刷新待处理的任务状态更新到数据库
        """
        if not self._pending_task_updates:
            return
        with self.app.app_context():
            for task_id, status in self._pending_task_updates.items():
                NewTaskModel.query.filter_by(id=task_id).update({"status": status})
            db.session.commit()
            self._pending_task_updates.clear()

    def pause_task(self, task_id):
        """
        暂停任务
        :param task_id: 任务ID
        :return:
        """
        # 与网络线程 check_execute_tasks 共用同一把锁（OCC 的 task_lock），保护 new_tasks/pause_tasks/net_tasks_buffer
        lock = self.operation_center.task_lock if self.operation_center else nullcontext()
        with lock:
            if task_id in self.net_tasks_buffer:
                main_task = self.net_tasks_buffer[task_id]
                if main_task.status == '正在执行':
                    # 正在执行的任务无法暂停
                    return False
                main_task.status = '暂停'
                for task in main_task.subtasks:
                    if task in self.new_tasks:
                        self.new_tasks.remove(task)
                    self.pause_tasks.append(task)
                self.update_task(main_task)
                return True
            else:
                for new_task in self.new_tasks:
                    if new_task.task_id == task_id and new_task.status == '等待执行':
                        self.pause_tasks.append(new_task)
                        self.new_tasks.remove(new_task)
                        new_task.status = '暂停'
                        self.update_task(new_task)
                        return True

    def start_task(self, task_id):
        """
        启动任务
        :param task_id: 任务ID
        :return:
        """
        # 与网络线程 check_execute_tasks 共用同一把锁（OCC 的 task_lock），保护 new_tasks/pause_tasks/net_tasks_buffer
        lock = self.operation_center.task_lock if self.operation_center else nullcontext()
        with lock:
            if task_id in self.net_tasks_buffer:
                main_task = self.net_tasks_buffer[task_id]
                if main_task.status == '暂停':
                    # 暂停的任务可以继续执行
                    main_task.status = '等待执行'
                    for task in main_task.subtasks:
                        if task in self.pause_tasks:
                            self.pause_tasks.remove(task)
                        self.new_tasks.append(task)
                    self.update_task(main_task)
                    return True
            else:
                for task in self.pause_tasks:
                    if task.task_id == task_id:
                        self.pause_tasks.remove(task)
                        task.status = '等待执行'
                        self.update_task(task)
                        self.new_tasks.append(task)
                        return True

    def run(self):
        """
        卫星网络线程运行函数
        """
        # 使用保存的app实例
        app = self.app

        with app.app_context() if app else nullcontext():
            print("卫星网络线程启动")
            # 启动socket服务器
            self._start_socket_server()
            # 主循环
            while True:  # 更新网络状态和收集结果
                try:
                    print("卫星网络的当前时间：", self.now_time)
                    # 非阻塞接收新客户端连接（无客户端时规划/仿真照常运行）
                    self._try_accept_client()
                    # 检查所有的卫星，执行任务
                    self.check_execute_tasks(self.now_time, self.time_multiple)
                    self.collect_results()
                    if self.is_planed:
                        self.update_net_state(self.now_time, self.time_multiple)
                    print("卫星网络中是否有任务：", self.new_tasks != [])
                    # self.collect_results(data_center_queue)
                    # 批量刷新任务状态更新
                    self.flush_task_updates()
                except Exception as e:
                    # 后台主循环异常兜底：打印异常并继续，避免线程退出
                    import traceback
                    print(f"卫星网络主循环发生异常: {e}")
                    print(traceback.format_exc())
                # 更新时间步
                time.sleep(INTER_VAL_TIME)  # 控制更新频率
